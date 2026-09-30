#!/usr/bin/env python3
"""Finish an openpyxl workbook as a macro-enabled .xlsm.

openpyxl writes cells, tables, formats, names and charts. This module adds what
it cannot write, by editing the saved package:

* table slicers (Class and KS2 band on every year sheet), mirroring the XML
  Excel itself wrote in the school's Year 11 tracker;
* buttons - rounded rectangles that run a macro when clicked;
* the VBA project, built from ./vba by vba_project.py;
* Excel's own spelling of line breaks in table column names (_x000a_);
* optionally, calculated values for every formula, computed by LibreOffice, so
  the file shows its numbers even before Excel recalculates (Protected View,
  previews on SharePoint and phones).
"""

import contextlib
import datetime as dt
import io
import os
import re
import shutil
import subprocess
import tempfile
import time
import uuid
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from lxml import etree

import vba_project

HERE = Path(__file__).resolve().parent
VBA_DIR = HERE / "vba"

NS = {
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "x14": "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main",
    "x15": "http://schemas.microsoft.com/office/spreadsheetml/2010/11/main",
}
REL_DRAWING = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing"
REL_SLICER = "http://schemas.microsoft.com/office/2007/relationships/slicer"
REL_SLICER_CACHE = "http://schemas.microsoft.com/office/2007/relationships/slicerCache"
REL_VBA = "http://schemas.microsoft.com/office/2006/relationships/vbaProject"
CT_DRAWING = "application/vnd.openxmlformats-officedocument.drawing+xml"
CT_SLICER = "application/vnd.ms-excel.slicer+xml"
CT_SLICER_CACHE = "application/vnd.ms-excel.slicerCache+xml"
CT_VBA = "application/vnd.ms-office.vbaProject"
CT_XLSX_MAIN = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
CT_XLSM_MAIN = "application/vnd.ms-excel.sheet.macroEnabled.main+xml"
EMU_PER_PX = 9525

BUTTON_FILL = "2A78D6"
BUTTON_INK = "FFFFFF"


# ---------------------------------------------------------------- extras model

@dataclass
class Shape:
    kind: str               # "button" or "slicer"
    name: str
    x: int                  # pixels from the sheet's top-left corner
    y: int
    w: int
    h: int
    macro: str = ""
    text: str = ""


@dataclass
class SlicerSpec:
    name: str               # shape / slicer name, unique in the workbook
    cache: str              # slicer cache (defined) name
    caption: str
    table: str
    column_id: int
    columns: int            # buttons per row
    column_name: str = ""   # the table column it filters (Excel's sourceName)


@dataclass
class SheetExtras:
    ws: object
    shapes: list = field(default_factory=list)
    slicers: list = field(default_factory=list)


class Extras:
    """What xlsm_package adds after openpyxl has saved the workbook."""

    def __init__(self):
        self.sheets = {}

    def _sheet(self, ws):
        return self.sheets.setdefault(ws.title, SheetExtras(ws))

    def add_buttons(self, ws, buttons, left_px, top_px, width_px=110, height_px=26, gap_px=8):
        sx = self._sheet(ws)
        x = left_px
        for label, macro in buttons:
            sx.shapes.append(Shape("button", f"btn{macro}_{len(sx.shapes) + 1}", x, top_px,
                                   width_px, height_px, macro=macro, text=label))
            x += width_px + gap_px

    def add_year_sheet(self, ws, table, class_col_id, band_col_id, col_widths, row_height, suffix):
        """Buttons and two slicers in the frozen corner above UPN and the names (A1:C1)."""
        sx = self._sheet(ws)
        corner_w = sum(col_px(w) for w in col_widths)
        pad = 6
        bw = (corner_w - 2 * pad - 2 * 6) // 3
        for i, (label, macro) in enumerate((("Manage tests", "ShowManageTests"),
                                            ("Dashboard", "GoToDashboard"), ("Start", "GoToStart"))):
            sx.shapes.append(Shape("button", f"btn{macro}_{suffix}", pad + i * (bw + 6), pad, bw, 24,
                                   macro=macro, text=label))
        top = pad + 24 + 8
        height = int(row_height * 96 / 72) - top - 40          # leave room for the header text
        half = (corner_w - 2 * pad - 6) // 2
        specs = [SlicerSpec(f"Class {suffix}", f"Slicer_Class_{suffix}", "Class", table, class_col_id, 2,
                            "Class"),
                 SlicerSpec(f"KS2 band {suffix}", f"Slicer_KS2_Band_{suffix}", "KS2 band", table, band_col_id, 3,
                            "Avg KS2 Band")]
        for i, spec in enumerate(specs):
            sx.slicers.append(spec)
            sx.shapes.append(Shape("slicer", spec.name, pad + i * (half + 6), top, half, height))


def col_px(width):
    """Excel column width (characters) -> pixels, for a 7-pixel maximum digit width."""
    return int(((256 * width + int(128 / 7)) / 256) * 7)


# --------------------------------------------------------------- xml helpers

def _q(prefix, tag):
    return f"{{{NS[prefix]}}}{tag}"


def _parse(data):
    return etree.fromstring(data)


def _dump(root):
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _stable_guid(*parts):
    return "{" + str(uuid.uuid5(uuid.NAMESPACE_URL, "|".join(parts))).upper() + "}"


class Package:
    def __init__(self, data: bytes):
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            self.parts = {n: z.read(n) for n in z.namelist()}

    def xml(self, name):
        return _parse(self.parts[name])

    def put(self, name, root_or_bytes):
        self.parts[name] = root_or_bytes if isinstance(root_or_bytes, bytes) else _dump(root_or_bytes)

    def rels_name(self, part):
        folder, base = part.rsplit("/", 1)
        return f"{folder}/_rels/{base}.rels"

    def rels(self, part):
        name = self.rels_name(part)
        if name in self.parts:
            return _parse(self.parts[name])
        return etree.Element(_q("rel", "Relationships"), nsmap={None: NS["rel"]})

    def add_rel(self, part, rel_type, target):
        root = self.rels(part)
        ids = {r.get("Id") for r in root}
        n = 1
        while f"rId{n}" in ids:
            n += 1
        el = etree.SubElement(root, _q("rel", "Relationship"))
        el.set("Id", f"rId{n}")
        el.set("Type", rel_type)
        el.set("Target", target)
        self.put(self.rels_name(part), root)
        return f"rId{n}"

    def content_type(self, part, ctype):
        root = self.xml("[Content_Types].xml")
        for o in root.findall(_q("ct", "Override")):
            if o.get("PartName") == "/" + part:
                o.set("ContentType", ctype)
                break
        else:
            el = etree.SubElement(root, _q("ct", "Override"))
            el.set("PartName", "/" + part)
            el.set("ContentType", ctype)
        self.put("[Content_Types].xml", root)

    def default_type(self, ext, ctype):
        root = self.xml("[Content_Types].xml")
        if not any(d.get("Extension") == ext for d in root.findall(_q("ct", "Default"))):
            el = etree.Element(_q("ct", "Default"))
            el.set("Extension", ext)
            el.set("ContentType", ctype)
            root.insert(0, el)
        self.put("[Content_Types].xml", root)

    def sheet_parts(self):
        """{sheet name: part name} from workbook.xml and its relationships."""
        wb = self.xml("xl/workbook.xml")
        rels = {r.get("Id"): r.get("Target") for r in self.rels("xl/workbook.xml")}
        out = {}
        for s in wb.find(_q("main", "sheets")):
            out[s.get("name")] = _resolve("xl/workbook.xml", rels[s.get(_q("r", "id"))])
        return out

    def table_ids(self):
        out = {}
        for name, data in self.parts.items():
            if name.startswith("xl/tables/") and name.endswith(".xml"):
                root = _parse(data)
                out[root.get("displayName")] = (int(root.get("id")), name)
        return out

    def tobytes(self):
        bio = io.BytesIO()
        order = ["[Content_Types].xml"] + sorted(n for n in self.parts if n != "[Content_Types].xml")
        with zipfile.ZipFile(bio, "w", zipfile.ZIP_DEFLATED) as z:
            for n in order:
                z.writestr(n, self.parts[n])
        return bio.getvalue()


# ------------------------------------------------------------------ drawings

def _anchor_cells(ws, x, y):
    """Pixel position -> (col, colOff EMU, row, rowOff EMU) using the sheet's sizes."""
    col, acc = 0, 0
    while True:
        letter = _col_letter(col + 1)
        dim = ws.column_dimensions.get(letter)
        w = col_px(dim.width) if dim is not None and dim.width else 64
        if acc + w > x:
            break
        acc += w
        col += 1
    col_off = (x - acc) * EMU_PER_PX
    row, acc = 0, 0
    while True:
        dim = ws.row_dimensions.get(row + 1)
        h = round((dim.height if dim is not None and dim.height else 15) * 96 / 72)
        if acc + h > y:
            break
        acc += h
        row += 1
    return col, col_off, row, (y - acc) * EMU_PER_PX


def _col_letter(n):
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _anchor_xml(ws, shape, shape_id):
    c1, o1, r1, p1 = _anchor_cells(ws, shape.x, shape.y)
    c2, o2, r2, p2 = _anchor_cells(ws, shape.x + shape.w, shape.y + shape.h)
    frm = (f"<xdr:from><xdr:col>{c1}</xdr:col><xdr:colOff>{o1}</xdr:colOff>"
           f"<xdr:row>{r1}</xdr:row><xdr:rowOff>{p1}</xdr:rowOff></xdr:from>"
           f"<xdr:to><xdr:col>{c2}</xdr:col><xdr:colOff>{o2}</xdr:colOff>"
           f"<xdr:row>{r2}</xdr:row><xdr:rowOff>{p2}</xdr:rowOff></xdr:to>")
    ext = f'<a:off x="{shape.x * EMU_PER_PX}" y="{shape.y * EMU_PER_PX}"/><a:ext cx="{shape.w * EMU_PER_PX}" cy="{shape.h * EMU_PER_PX}"/>'
    if shape.kind == "button":
        body = (
            f'<xdr:sp macro="[0]!{shape.macro}" textlink="">'
            f'<xdr:nvSpPr><xdr:cNvPr id="{shape_id}" name="{shape.name}" descr="Runs {shape.macro}"/>'
            f'<xdr:cNvSpPr/></xdr:nvSpPr>'
            f'<xdr:spPr><a:xfrm>{ext}</a:xfrm><a:prstGeom prst="roundRect"><a:avLst/></a:prstGeom>'
            f'<a:solidFill><a:srgbClr val="{BUTTON_FILL}"/></a:solidFill><a:ln><a:noFill/></a:ln></xdr:spPr>'
            f'<xdr:txBody><a:bodyPr vertOverflow="clip" horzOverflow="clip" wrap="none" lIns="36000" '
            f'rIns="36000" tIns="0" bIns="0" rtlCol="0" anchor="ctr"/><a:lstStyle/>'
            f'<a:p><a:pPr algn="ctr"/><a:r><a:rPr lang="en-GB" sz="1000" b="1">'
            f'<a:solidFill><a:srgbClr val="{BUTTON_INK}"/></a:solidFill><a:latin typeface="Arial"/>'
            f'</a:rPr><a:t>{shape.text}</a:t></a:r></a:p></xdr:txBody></xdr:sp>')
    else:
        body = (
            f'<mc:AlternateContent xmlns:mc="{NS["mc"]}" '
            f'xmlns:sle15="http://schemas.microsoft.com/office/drawing/2012/slicer">'
            f'<mc:Choice Requires="sle15"><xdr:graphicFrame macro="">'
            f'<xdr:nvGraphicFramePr><xdr:cNvPr id="{shape_id}" name="{shape.name}"/><xdr:cNvGraphicFramePr/>'
            f'</xdr:nvGraphicFramePr><xdr:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/></xdr:xfrm>'
            f'<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/drawing/2010/slicer">'
            f'<sle:slicer xmlns:sle="http://schemas.microsoft.com/office/drawing/2010/slicer" name="{shape.name}"/>'
            f'</a:graphicData></a:graphic></xdr:graphicFrame></mc:Choice>'
            f'<mc:Fallback xmlns=""><xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="0" name=""/>'
            f'<xdr:cNvSpPr><a:spLocks noTextEdit="1"/></xdr:cNvSpPr></xdr:nvSpPr><xdr:spPr><a:xfrm>{ext}</a:xfrm>'
            f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill><a:prstClr val="white"/></a:solidFill>'
            f'<a:ln w="1"><a:solidFill><a:prstClr val="green"/></a:solidFill></a:ln></xdr:spPr>'
            f'<xdr:txBody><a:bodyPr vertOverflow="clip" horzOverflow="clip"/><a:lstStyle/><a:p><a:r>'
            f'<a:rPr lang="en-GB" sz="1100"/><a:t>This shape represents a table slicer. Table slicers need '
            f'Excel 2013 or later.</a:t></a:r></a:p></xdr:txBody></xdr:sp></mc:Fallback></mc:AlternateContent>')
    xml = (f'<xdr:twoCellAnchor xmlns:xdr="{NS["xdr"]}" xmlns:a="{NS["a"]}" editAs="absolute">'
           f'{frm}{body}<xdr:clientData/></xdr:twoCellAnchor>')
    return etree.fromstring(xml)


def _drawing_for_sheet(pkg, sheet_part):
    """Existing drawing part of a sheet, or a new empty one wired into the sheet."""
    rels = pkg.rels(sheet_part)
    for r in rels:
        if r.get("Type") == REL_DRAWING:
            return _resolve(sheet_part, r.get("Target"))
    n = 1
    while f"xl/drawings/drawing{n}.xml" in pkg.parts:
        n += 1
    path = f"xl/drawings/drawing{n}.xml"
    root = etree.Element(_q("xdr", "wsDr"), nsmap={"xdr": NS["xdr"], "a": NS["a"]})
    pkg.put(path, root)
    pkg.content_type(path, CT_DRAWING)
    rid = pkg.add_rel(sheet_part, REL_DRAWING, f"../drawings/drawing{n}.xml")
    sheet = pkg.xml(sheet_part)
    el = etree.Element(_q("main", "drawing"))
    el.set(_q("r", "id"), rid)
    _insert_before(sheet, el, ["legacyDrawing", "legacyDrawingHF", "drawingHF", "picture", "oleObjects",
                               "controls", "webPublishItems", "tableParts", "extLst"])
    pkg.put(sheet_part, sheet)
    return path


def _resolve(source_part, target):
    """A relationship target (relative or absolute) -> a part name."""
    if target.startswith("/"):
        return target[1:]
    return os.path.normpath(os.path.join(os.path.dirname(source_part), target)).replace("\\", "/")


def _insert_before(root, el, later_tags):
    for i, child in enumerate(root):
        if etree.QName(child).localname in later_tags:
            root.insert(i, el)
            return
    root.append(el)


# ------------------------------------------------------------------- slicers

SLICER_HEAD = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
               f'<slicers xmlns="{NS["x14"]}" xmlns:mc="{NS["mc"]}" mc:Ignorable="x xr10" '
               f'xmlns:x="{NS["main"]}" '
               'xmlns:xr10="http://schemas.microsoft.com/office/spreadsheetml/2016/revision10">')


def _add_slicers(pkg, sheet_name, sheet_part, specs, table_ids, counter):
    cache_rids = []
    n_slicer = 1
    while f"xl/slicers/slicer{n_slicer}.xml" in pkg.parts:
        n_slicer += 1
    body = []
    for spec in specs:
        counter[0] += 1
        n = counter[0]
        table_id, table_part = table_ids[spec.table]
        # x15:tableSlicerCache/@column is the table column's id, not its position.
        ids = {tc.get("name"): tc.get("id") for tc in pkg.xml(table_part).iter(f"{{{NS['main']}}}tableColumn")}
        if spec.column_name:
            if spec.column_name not in ids:
                raise ValueError(f"slicer {spec.name}: no column {spec.column_name!r} in {spec.table}")
            spec.column_id = int(ids[spec.column_name])
        cache_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            f'<slicerCacheDefinition xmlns="{NS["x14"]}" xmlns:mc="{NS["mc"]}" mc:Ignorable="x xr10" '
            f'xmlns:x="{NS["main"]}" xmlns:xr10="http://schemas.microsoft.com/office/spreadsheetml/2016/revision10" '
            f'name="{spec.cache}" xr10:uid="{_stable_guid("cache", spec.cache)}" '
            f'sourceName="{spec.column_name or spec.caption}">'
            f'<extLst><x:ext uri="{{2F2917AC-EB37-4324-AD4E-5DD8C200BD13}}" xmlns:x15="{NS["x15"]}">'
            f'<x15:tableSlicerCache tableId="{table_id}" column="{spec.column_id}"/></x:ext></extLst>'
            '</slicerCacheDefinition>')
        cache_part = f"xl/slicerCaches/slicerCache{n}.xml"
        pkg.put(cache_part, cache_xml.encode())
        pkg.content_type(cache_part, CT_SLICER_CACHE)
        cache_rids.append((spec.cache, pkg.add_rel("xl/workbook.xml", REL_SLICER_CACHE,
                                                   f"slicerCaches/slicerCache{n}.xml")))
        body.append(f'<slicer name="{spec.name}" xr10:uid="{_stable_guid("slicer", spec.name)}" '
                    f'cache="{spec.cache}" caption="{spec.caption}" columnCount="{spec.columns}" '
                    f'rowHeight="180000"/>')
    slicer_part = f"xl/slicers/slicer{n_slicer}.xml"
    pkg.put(slicer_part, (SLICER_HEAD + "".join(body) + "</slicers>").encode())
    pkg.content_type(slicer_part, CT_SLICER)
    rid = pkg.add_rel(sheet_part, REL_SLICER, f"../slicers/slicer{n_slicer}.xml")
    sheet = pkg.xml(sheet_part)
    ext_lst = sheet.find(_q("main", "extLst"))
    if ext_lst is None:
        ext_lst = etree.SubElement(sheet, _q("main", "extLst"))
    ext = etree.SubElement(ext_lst, _q("main", "ext"), nsmap={"x15": NS["x15"]})
    ext.set("uri", "{3A4CF648-6AED-40f4-86FF-DC5316D8AED3}")
    lst = etree.SubElement(ext, _q("x14", "slicerList"), nsmap={"x14": NS["x14"]})
    sl = etree.SubElement(lst, _q("x14", "slicer"))
    sl.set(_q("r", "id"), rid)
    pkg.put(sheet_part, sheet)
    return cache_rids


def _register_slicer_caches(pkg, cache_rids):
    wb = pkg.xml("xl/workbook.xml")
    names = wb.find(_q("main", "definedNames"))
    if names is None:
        names = etree.Element(_q("main", "definedNames"))
        _insert_before(wb, names, ["calcPr", "oleSize", "customWorkbookViews", "pivotCaches",
                                   "smartTagPr", "smartTagTypes", "webPublishing", "fileRecoveryPr",
                                   "webPublishObjects", "extLst"])
    for cache, _ in cache_rids:
        dn = etree.SubElement(names, _q("main", "definedName"))
        dn.set("name", cache)
        dn.text = "#N/A"
    ext_lst = wb.find(_q("main", "extLst"))
    if ext_lst is None:
        ext_lst = etree.SubElement(wb, _q("main", "extLst"))
    ext = etree.SubElement(ext_lst, _q("main", "ext"), nsmap={"x15": NS["x15"]})
    ext.set("uri", "{46BE6895-7355-4a93-B00E-2C351335B9C9}")
    caches = etree.SubElement(ext, _q("x15", "slicerCaches"), nsmap={"x14": NS["x14"]})
    for _, rid in cache_rids:
        el = etree.SubElement(caches, _q("x14", "slicerCache"))
        el.set(_q("r", "id"), rid)
    pkg.put("xl/workbook.xml", wb)


# ----------------------------------------------------------------- VBA project

def vba_modules(sheet_codenames):
    """The project's modules: ThisWorkbook, one per sheet, the code modules, the form."""
    def src(name):
        return (VBA_DIR / name).read_text(encoding="utf-8")

    sheet_code = {"shDashboard": src("shDashboard.vba")}
    mods = [vba_project.Module("ThisWorkbook", "document", src("ThisWorkbook.vba"), base=vba_project.WORKBOOK_BASE)]
    for code in sheet_codenames:
        mods.append(vba_project.Module(code, "document", sheet_code.get(code, ""),
                                       base=vba_project.WORKSHEET_BASE))
    mods.append(vba_project.Module("modTracker", "standard", src("modTracker.bas")))
    mods.append(vba_project.Module("modImport", "standard", src("modImport.bas")))
    mods.append(vba_project.Module("frmTests", "form", src("frmTests.vba"), caption="Manage tests",
                                   width_pt=392, height_pt=312))
    return mods


def _add_vba(pkg, wb):
    codenames = [ws.sheet_properties.codeName for ws in wb.worksheets]
    data = vba_project.build_vba_project(vba_modules(codenames), seed="mastery-quiz-tracker")
    pkg.put("xl/vbaProject.bin", data)
    pkg.default_type("bin", CT_VBA)
    pkg.content_type("xl/workbook.xml", CT_XLSM_MAIN)
    pkg.add_rel("xl/workbook.xml", REL_VBA, "vbaProject.bin")


# -------------------------------------------------------------- table columns

def _chart_number_formats(pkg):
    """Data labels keep their own number format rather than the source cells'."""
    for name in list(pkg.parts):
        if name.startswith("xl/charts/chart") and name.endswith(".xml"):
            text = pkg.parts[name].decode("utf-8")
            text = re.sub(r'<(\w+:)?numFmt formatCode="([^"]*)"\s*/>',
                          lambda m: f'<{m.group(1) or ""}numFmt formatCode="{m.group(2)}" sourceLinked="0"/>', text)
            pkg.parts[name] = text.encode("utf-8")


def _excel_line_breaks(pkg):
    """Excel writes a line break in a table column name as _x000a_."""
    for name in list(pkg.parts):
        if name.startswith("xl/tables/") and name.endswith(".xml"):
            text = pkg.parts[name].decode("utf-8")
            text = re.sub(r'(<tableColumn [^>]*name=")([^"]*)(")',
                          lambda m: m.group(1) + m.group(2).replace("&#10;", "_x000a_").replace("\n", "_x000a_")
                          + m.group(3), text)
            pkg.parts[name] = text.encode("utf-8")


# --------------------------------------------------------------- cached values

def uno_prop(name, value):
    from com.sun.star.beans import PropertyValue
    p = PropertyValue()
    p.Name, p.Value = name, value
    return p


@contextlib.contextmanager
def libreoffice():
    """A headless LibreOffice process; yields its Desktop object."""
    import uno  # noqa: F401 - system LibreOffice Python bindings

    profile = tempfile.mkdtemp(prefix="lo-profile-")
    pipe = f"mqt_{os.getpid()}_{int(time.time() * 1000)}"
    proc = subprocess.Popen(
        ["soffice", "--headless", "--norestore", "--nologo", "--nodefault",
         f"-env:UserInstallation=file://{profile}", f"--accept=pipe,name={pipe};urp;StarOffice.ComponentContext"],
        stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    try:
        local = uno.getComponentContext()
        resolver = local.ServiceManager.createInstanceWithContext("com.sun.star.bridge.UnoUrlResolver", local)
        deadline = time.time() + 60
        while True:
            try:
                ctx = resolver.resolve(f"uno:pipe,name={pipe};urp;StarOffice.ComponentContext")
                break
            except Exception:
                if time.time() > deadline:
                    raise RuntimeError("LibreOffice did not start")
                time.sleep(0.3)
        desktop = ctx.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", ctx)
        yield desktop
        try:
            desktop.terminate()
        except Exception:
            pass
    finally:
        try:
            proc.wait(timeout=60)
        except subprocess.TimeoutExpired:
            proc.kill()
        shutil.rmtree(profile, ignore_errors=True)


def open_document(desktop, path, run_macros=False):
    import uno
    return desktop.loadComponentFromURL(uno.systemPathToFileUrl(str(Path(path).resolve())), "_blank", 0,
                                        (uno_prop("Hidden", True), uno_prop("MacroExecutionMode", 4 if run_macros else 0)))


def recalculate(pairs, timeout=900):
    """Open each source workbook in headless LibreOffice, recalculate, save an .xlsx copy.

    pairs: [(source path, destination path)] - one LibreOffice session serves them all.
    """
    import uno
    with libreoffice() as desktop:
        for src, dst in pairs:
            doc = open_document(desktop, src)
            doc.calculateAll()
            doc.storeToURL(uno.systemPathToFileUrl(str(Path(dst).resolve())),
                           (uno_prop("FilterName", "Calc MS Excel 2007 XML"),))
            doc.close(True)


def recalculate_copy(src: Path, dst: Path, timeout=900):
    recalculate([(src, dst)], timeout)


ERRORS = {"#N/A", "#VALUE!", "#REF!", "#DIV/0!", "#NUM!", "#NAME?", "#NULL!"}
EPOCH = dt.datetime(1899, 12, 30)


def cell_values(xlsx_path):
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path, data_only=True)
    out = {}
    for ws in wb.worksheets:
        vals = {}
        for row in ws.iter_rows():
            for c in row:
                vals[c.coordinate] = c.value
        out[ws.title] = vals
    return out


def _fill_cached_values(pkg, values):
    for sheet, part in pkg.sheet_parts().items():
        if sheet not in values:
            continue
        root = pkg.xml(part)
        vals = values[sheet]
        for c in root.iter(_q("main", "c")):
            f = c.find(_q("main", "f"))
            if f is None:
                continue
            for old in c.findall(_q("main", "v")):
                c.remove(old)
            v = vals.get(c.get("r"))
            ve = etree.SubElement(c, _q("main", "v"))
            if isinstance(v, bool):
                c.set("t", "b")
                ve.text = "1" if v else "0"
            elif isinstance(v, (int, float)):
                c.attrib.pop("t", None)
                ve.text = repr(float(v)) if isinstance(v, float) else str(v)
            elif isinstance(v, dt.datetime):
                c.attrib.pop("t", None)
                ve.text = repr((v - EPOCH).total_seconds() / 86400)
            elif isinstance(v, dt.date):
                c.attrib.pop("t", None)
                ve.text = str((dt.datetime(v.year, v.month, v.day) - EPOCH).days)
            elif isinstance(v, str) and v in ERRORS:
                c.set("t", "e")
                ve.text = v
            else:
                c.set("t", "str")
                ve.text = "" if v is None else str(v)
        pkg.put(part, root)


# ---------------------------------------------------------------------- save

def assemble(wb, extras: Extras) -> bytes:
    bio = io.BytesIO()
    wb.save(bio)
    pkg = Package(bio.getvalue())
    _excel_line_breaks(pkg)
    _chart_number_formats(pkg)
    sheet_parts = pkg.sheet_parts()
    table_ids = pkg.table_ids()
    counter = [0]
    all_caches = []
    for title, sx in extras.sheets.items():
        part = sheet_parts[title]
        if sx.slicers:
            all_caches += _add_slicers(pkg, title, part, sx.slicers, table_ids, counter)
        if sx.shapes:
            drawing = _drawing_for_sheet(pkg, part)
            root = pkg.xml(drawing)
            next_id = 1000 + len(root)
            for shape in sx.shapes:
                root.append(_anchor_xml(sx.ws, shape, next_id))
                next_id += 1
            pkg.put(drawing, root)
    if all_caches:
        _register_slicer_caches(pkg, all_caches)
    _add_vba(pkg, wb)
    return pkg.tobytes()


def save(wb, extras: Extras, out: Path, fill_cache=True):
    data = assemble(wb, extras)
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    if fill_cache and shutil.which("soffice"):
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "book.xlsm"
            dst = Path(tmp) / "book-calc.xlsx"
            src.write_bytes(data)
            recalculate_copy(src, dst)
            values = cell_values(dst)
        pkg = Package(data)
        _fill_cached_values(pkg, values)
        data = pkg.tobytes()
    elif fill_cache:
        print("LibreOffice not found: saved without cached values (Excel calculates them on opening).")
    out.write_bytes(data)
    return out
