#!/usr/bin/env python3
"""Generate the Science Mastery Quiz Tracker (Years 7-11) as a macro-enabled workbook.

One workbook holds every year group:

* Year 7 ... Year 11 - one Excel table per year group. Each test adds exactly two
  columns (Raw Score and Stanine); there are no empty placeholder columns.
* Assessment Info    - the register of tests and the stanine lookup.
* Overview           - how every class in every year group is doing.
* Dashboard          - one year group: each test for the year and each class,
                       a test drill-down, stanine spread, group gaps and a trend.
* Watch List         - students working well below their KS2 starting point.
* Start / Settings   - instructions, the stanine key, class teachers, rules.

The Manage tests form (add / edit / remove), the student update from a new
MIS export and the navigation buttons are VBA in ./vba. vba_project.py turns
that source into the workbook's vbaProject.bin, and xlsm_package.py adds the
parts openpyxl cannot write (slicers, buttons, the VBA project).

Usage:
  python3 generate_tracker.py                       # demo with fictional students
  python3 generate_tracker.py --students private/All_Students.xlsx \\
      --carry-over "Year 11 tracker.xlsx" --out "Science Mastery Quiz Tracker 2026-27.xlsm"

Real student data never belongs in this (public) repository: write real
workbooks outside it, or into the git-ignored private/ folder.
"""

import argparse
import datetime as dt
import math
import random
import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart
from openpyxl.chart.data_source import AxDataSource, NumDataSource, NumRef, StrRef
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint, Marker
from openpyxl.chart.series import Series, SeriesLabel
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, Font as DFont
from openpyxl.formatting.rule import DataBarRule, FormatObject, FormulaRule, IconSet, Rule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.styles.numbers import NumberFormat
from openpyxl.utils import get_column_letter
from openpyxl.utils.indexed_list import IndexedList
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableColumn, TableFormula, TableStyleInfo

import xlsm_package

HERE = Path(__file__).resolve().parent
VBA_DIR = HERE / "vba"

# ------------------------------------------------------------ shared with VBA


def read_vba_constants(path=VBA_DIR / "modTracker.bas"):
    """Public Const values from modTracker.bas, so the macro and the generator agree."""
    consts = {}
    for m in re.finditer(r'^Public Const (\w+) As (\w+) = (.+)$', path.read_text(), re.M):
        name, kind, raw = m.groups()
        raw = raw.split("'")[0].strip() if not raw.startswith('"') else raw.strip()
        if kind == "String":
            consts[name] = raw[1:-1].replace('""', '"')
        elif raw.startswith("&H"):
            bgr = int(raw[2:].rstrip("&"), 16)
            consts[name] = f"{bgr & 0xFF:02X}{(bgr >> 8) & 0xFF:02X}{(bgr >> 16) & 0xFF:02X}"
        else:
            consts[name] = float(raw) if "." in raw else int(raw)
    return consts


VBA = read_vba_constants()
RAW_SUFFIX = VBA["RAW_SUFFIX"]
STANINE_SUFFIX = VBA["STANINE_SUFFIX"]
FIRST_TEST_COL = VBA["FIRST_TEST_COL"]

# ------------------------------------------------------------------- layout

YEARS = [7, 8, 9, 10, 11]
YEAR_NAME = {y: f"Year {y}" for y in YEARS}
TABLE = {y: f"tblY{y}" for y in YEARS}
CODENAME = {y: f"shY{y}" for y in YEARS}
REG = "tblAssessments"
TEACHERS = "tblTeachers"

FIXED = [  # (header, width, number format)
    ("UPN", 15.5, "@"),
    ("Preferred Last name", 17, "@"),
    ("Preferred First name", 15, "@"),
    ("Sex Code", 5.3, "General"),
    ("SEN Status Code", 5.3, "General"),
    ("PP Deprivation", 5.3, "General"),
    ("Att%", 6.3, "0.0"),
    ("Avg KS2", 7, "0.0"),
    ("Avg KS2 Band", 6.3, "0"),
    ("Class", 15.5, "General"),
    ("Tests Sat", 5.8, "0"),
    ("Mean Stanine", 6.8, "0.0"),
    ("vs KS2 Band", 6.8, "+0.0;-0.0;0.0"),
]
FORMULA_COLS = {"Avg KS2 Band", "Tests Sat", "Mean Stanine", "vs KS2 Band"}
assert len(FIXED) + 1 == FIRST_TEST_COL, "fixed columns must end just before FIRST_TEST_COL"
FIRST_TEST = get_column_letter(FIRST_TEST_COL)          # N
LAST_TEST = "ZZ"                                        # room for 344 tests per year

MAX_STUDENTS = 600       # helper rows per year on the Calc sheet
MAX_TESTS = 60           # tests per year shown on the Dashboard
CLASS_SLOTS = 8          # classes per year shown on the dashboards
WATCH_ROWS = 50

STANINE_Z = [-99, -1.75, -1.25, -0.75, -0.25, 0.25, 0.75, 1.25, 1.75]
EXPECTED_PCT = [4, 7, 12, 17, 20, 17, 12, 7, 4]
STANINE_WORDS = ["Very low", "Low", "Below average", "Slightly below average", "Average",
                 "Slightly above average", "Above average", "High", "Very high"]

MEASURES = ["Mean %", "Mean stanine", "Completion %"]

# ------------------------------------------------------------------ styling
# Colours follow a validated data-viz palette: blue for "above", red for
# "below" (a colour-blind-safe diverging pair), status colours only for
# completion, and every coloured cell also shows its number.

INK, INK2, MUTED = "0B0B0B", "52514E", "898781"
NAVY, BLUE, BLUE_500 = "0D366B", "2A78D6", "256ABF"
BLUE_100, BLUE_200 = "CDE2FB", "9EC5F4"
RED_100, RED_200 = "FBE4E3", "F2C4C3"
AMBER_100 = "FDEBC4"
GREY_FILL, GRID, BASE = "F5F5F4", "E1E0D9", "C3C2B7"
INPUT_FILL = "FFF3BF"
WHITE = "FFFFFF"
FONT = "Arial"
TEST_FILLS = (VBA["TEST_FILL_ODD"], VBA["TEST_FILL_EVEN"])
STANINE_FILL = VBA["STANINE_FILL"]


def font(size=10, bold=False, color=INK, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def fill(color):
    return PatternFill("solid", fgColor=color)


HAIR = Side(style="thin", color=GRID)
BOX = Border(left=HAIR, right=HAIR, top=HAIR, bottom=HAIR)
TOP_RULE = Border(top=Side(style="thin", color=BASE))
YEAR_RULE = Border(bottom=Side(style="thin", color=BASE))
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)
UNLOCKED = Protection(locked=False)


def put(ws, ref, value=None, *, f=None, fl=None, al=None, fmt=None, border=None):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = f or font()
    if fl:
        c.fill = fill(fl) if isinstance(fl, str) else fl
    if al:
        c.alignment = al
    if fmt:
        c.number_format = fmt
    if border:
        c.border = border
    return c


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def title_block(ws, title, subtitle, last_col):
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 30
    for col in range(1, last_col + 1):
        ws.cell(1, col).fill = fill(NAVY)
    put(ws, "B1", title, f=font(16, True, WHITE), fl=NAVY, al=LEFT)
    put(ws, "B2", subtitle, f=font(10, False, INK2, italic=True), al=LEFT)


def section(ws, ref, text):
    put(ws, ref, text, f=font(12, True, NAVY), al=LEFT)


def header_cells(ws, row, first_col, titles, fl=NAVY, color=WHITE, height=30):
    ws.row_dimensions[row].height = height
    for i, t in enumerate(titles):
        c = ws.cell(row, first_col + i, t)
        c.font = font(9, True, color)
        c.fill = fill(fl)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX


def selector(ws, ref, value, merge_to=None):
    c = put(ws, ref, value, f=font(11, True), fl=INPUT_FILL, al=LEFT, border=BOX)
    c.protection = UNLOCKED
    if merge_to:
        ws.merge_cells(f"{ref}:{merge_to}")
    return c


# ----------------------------------------------------------------- formulas

def tr(table, col):
    """A this-row structured reference, as Excel stores it in the file."""
    return f"{table}[[#This Row],[{col}]]"


def choose(index_expr, part):
    """CHOOSE(i, tblY7[part], ..., tblY11[part]) - pick a year group's table."""
    return f"CHOOSE({index_expr}," + ",".join(f"{TABLE[y]}[{part}]" for y in YEARS) + ")"


def to_file_formula(ui_formula):
    """A formula typed into Excel (as the VBA writes it) -> the text stored in the file."""
    body = ui_formula[1:] if ui_formula.startswith("=") else ui_formula
    return re.sub(r"(?<![\w.])STDEV\.P\(", "_xlfn.STDEV.P(", body)


def raw_header(test_name):
    return f"{test_name}\n{RAW_SUFFIX}"


def stanine_header(test_name):
    return f"{test_name}\n{STANINE_SUFFIX}"


def stanine_formula(table, raw_name):
    return to_file_formula(VBA["STANINE_FORMULA"].replace("{T}", table).replace("{R}", raw_name))


def raw_rule(first_cell, max_marks):
    return to_file_formula(VBA["RAW_RULE"].replace("{C}", first_cell).replace("{M}", f"{max_marks:g}"))


LF = 'CHAR(10)'


def Q(text):
    """A string literal inside a formula."""
    return '"' + text.replace('"', '""') + '"'


def ks2_band_formula(t):
    return (f'IF(ISNUMBER({tr(t, "Avg KS2")}),IFERROR(VLOOKUP(STANDARDIZE({tr(t, "Avg KS2")},'
            f'AVERAGE({t}[[Avg KS2]]),_xlfn.STDEV.P({t}[[Avg KS2]])),_Stanine,2),5),"")')


def tests_sat_formula(t, row):
    return (f'IF({tr(t, "Preferred Last name")}="","",COUNTIFS(${FIRST_TEST}$1:${LAST_TEST}$1,'
            f'"*"&{LF}&"{STANINE_SUFFIX}",${FIRST_TEST}{row}:${LAST_TEST}{row},">=1"))')


def mean_stanine_formula(row):
    return (f'IFERROR(AVERAGEIF(${FIRST_TEST}$1:${LAST_TEST}$1,"*"&{LF}&"{STANINE_SUFFIX}",'
            f'${FIRST_TEST}{row}:${LAST_TEST}{row}),"")')


def vs_ks2_formula(t):
    return (f'IF(AND(ISNUMBER({tr(t, "Mean Stanine")}),ISNUMBER({tr(t, "Avg KS2 Band")})),'
            f'{tr(t, "Mean Stanine")}-{tr(t, "Avg KS2 Band")},"")')


# --------------------------------------------------------------------- data

@dataclass
class Student:
    upn: str
    last: str
    first: str
    year: int
    sex: str = ""
    sen: str = ""
    pp: str = ""
    att: object = None
    ks2: object = "No Data"
    cls: str = ""

    @property
    def key(self):
        return self.upn.upper() if self.upn else f"{self.last}|{self.first}".upper()


@dataclass
class Test:
    year: int
    code: str
    title: str
    max_marks: float
    date: object = None
    scores: dict = field(default_factory=dict)      # Student.key -> number or "A"

    @property
    def name(self):
        return f"{self.code} - {self.title}" if self.code and self.title else (self.code or self.title)


@dataclass
class Dataset:
    students: list
    tests: list
    academic_year: str
    school: str
    demo: bool
    dashboard_year: int = 7
    watch_year: int = 7
    teachers: dict = field(default_factory=dict)       # (year, class) -> teacher
    dashboard_test: object = None                      # default: the year's latest test
    dashboard_class: object = None                     # default: the year's first class
    dashboard_measure: str = "Mean %"


def ks2_average(maths, reading):
    vals = [v for v in (maths, reading) if isinstance(v, (int, float)) and not isinstance(v, bool)]
    if not vals:
        return "No Data"
    return sum(vals) / len(vals)


def _header_index(headers, *patterns):
    for i, h in enumerate(headers):
        h = str(h or "").strip().lower()
        if h and any(re.fullmatch(p, h) for p in patterns):
            return i
    return None


def load_student_export(path):
    """Students in Years 7-11 from the MIS 'All Students' export (first sheet)."""
    ws = load_workbook(path, data_only=True, read_only=True).worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    hdr = next(i for i, r in enumerate(rows[:15]) if any(str(v or "").strip().upper() == "UPN" for v in r))
    headers = rows[hdr]
    col = {
        "upn": _header_index(headers, r"upn"),
        "last": _header_index(headers, r".*last name.*", r".*surname.*"),
        "first": _header_index(headers, r".*first name.*", r".*forename.*"),
        "year": _header_index(headers, r"year group.*", r"year"),
        "cls": _header_index(headers, r".*class.*", r".*teaching group.*"),
        "att": _header_index(headers, r".*attendance.*", r"att.*%.*"),
        "pp": _header_index(headers, r"pp.*", r".*pupil premium.*", r".*deprivation.*"),
        "sen": _header_index(headers, r"sen.*"),
        "sex": _header_index(headers, r"sex.*", r"gender.*"),
        "maths": _header_index(headers, r"ks2 maths.*"),
        "reading": _header_index(headers, r"ks2 reading.*"),
    }

    def get(row, key):
        i = col[key]
        return row[i] if i is not None and i < len(row) else None

    def text(v):
        return "" if v is None else str(v).strip()

    students = []
    for row in rows[hdr + 1:]:
        last = text(get(row, "last"))
        digits = re.sub(r"\D", "", text(get(row, "year")))
        if not last or last.lower().startswith("count:") or not digits or int(digits) not in YEARS:
            continue
        pp = text(get(row, "pp")).upper()
        pp = {"YES": "Y", "NO": "N", "TRUE": "Y", "FALSE": "N"}.get(pp, pp)
        att = get(row, "att")
        students.append(Student(
            upn=text(get(row, "upn")), last=last, first=text(get(row, "first")), year=int(digits),
            sex=text(get(row, "sex")).upper(), sen=text(get(row, "sen")).upper(), pp=pp,
            att=att if isinstance(att, (int, float)) and not isinstance(att, bool) else None,
            ks2=ks2_average(get(row, "maths"), get(row, "reading")), cls=text(get(row, "cls"))))
    return students


def load_carry_over(path, students):
    """Tests and scores from an existing single-year tracker (like the Year 11 file).

    Raw-score columns are headed "<test name>\\nRaw Score". Max marks come from its
    Assessment Info table, matched in order. Scores are matched to students by UPN,
    then by name; the year group is the one most matched students belong to.
    """
    wb = load_workbook(path, data_only=True)
    data_ws = next(ws for ws in wb.worksheets
                   if any(str(c.value or "").endswith("\n" + RAW_SUFFIX) for c in ws[1]))
    headers = [str(c.value or "") for c in data_ws[1]]
    upn_col = headers.index("UPN")
    last_col, first_col = headers.index("Preferred Last name"), headers.index("Preferred First name")
    max_marks = []
    info = next((ws for ws in wb.worksheets if ws.title.lower().startswith("assessment")), None)
    if info is not None:
        for r in range(2, info.max_row + 1):
            v = info.cell(r, 2).value
            if isinstance(v, (int, float)) and info.cell(r, 1).value:
                max_marks.append(v)
    by_upn = {s.upn.upper(): s for s in students if s.upn}
    by_name = {f"{s.last}|{s.first}".upper(): s for s in students}
    tests = []
    raw_cols = [i for i, h in enumerate(headers) if h.endswith("\n" + RAW_SUFFIX)]
    for n, ci in enumerate(raw_cols):
        name = headers[ci][: -len(RAW_SUFFIX) - 1].strip()
        code, _, title = name.partition(" - ")
        if not title:
            code, title = "", name
        scores, years = {}, []
        for row in data_ws.iter_rows(min_row=2, values_only=True):
            v = row[ci]
            if v is None or v == "":
                continue
            upn = str(row[upn_col] or "").strip().upper()
            s = by_upn.get(upn) or by_name.get(f"{row[last_col]}|{row[first_col]}".upper())
            if s is None:
                continue
            scores[s.key] = v if isinstance(v, (int, float)) else str(v).strip().upper()
            years.append(s.year)
        if not scores:
            continue
        year = max(set(years), key=years.count)
        best = max(v for v in scores.values() if isinstance(v, (int, float)))
        mm = max_marks[n] if n < len(max_marks) else max(best, 1)
        tests.append(Test(year, code.strip(), title.strip(), mm, None, scores))
    return tests


# Fictional data for the demo workbook ------------------------------------------------

FIRST_NAMES = ["Amelia", "Oliver", "Isla", "Noah", "Ava", "Leo", "Mia", "Arthur", "Ivy", "Theo",
               "Freya", "Harry", "Grace", "Oscar", "Lily", "Jack", "Rosie", "George", "Evie",
               "Alfie", "Sophie", "Charlie", "Ella", "Jacob", "Poppy", "Muhammad", "Aisha",
               "Zara", "Ibrahim", "Maya", "Kai", "Nia", "Tomasz", "Zofia", "Chidi", "Amara",
               "Rhys", "Seren", "Callum", "Niamh", "Ethan", "Hannah", "Dylan", "Megan", "Finley",
               "Layla", "Reuben", "Hollie", "Sami", "Priya", "Arjun", "Keira", "Lewis", "Maisie"]
LAST_NAMES = ["Smith", "Jones", "Taylor", "Brown", "Williams", "Wilson", "Johnson", "Davies",
              "Patel", "Robinson", "Wright", "Thompson", "Evans", "Walker", "White", "Roberts",
              "Green", "Hall", "Wood", "Jackson", "Clarke", "Khan", "Hughes", "Lewis", "Edwards",
              "Nowak", "Okafor", "Ahmed", "Murphy", "Kelly", "Price", "Bennett", "Griffiths",
              "Lloyd", "Morgan", "Shaw", "Fletcher", "Barker", "Pearson", "Hussain", "Begum",
              "Mistry", "Kowalski", "Adeyemi", "Harrison", "Cooper", "Ward", "Turner", "Parker"]

DEMO_CLASSES = {
    7: {"7A/Sc1": 30, "7A/Sc2": 30, "7B/Sc1": 27, "7B/Sc2": 26, "7B/Sc3": 18},
    8: {"8A/Sc1": 28, "8A/Sc2": 25, "8A/Sc3": 22, "8A/Sc4": 16, "Science Support": 2},
    9: {"9A/Sc1": 22, "9A/Sc2": 25, "9A/Sc3": 19, "9B/Sc1": 27, "9B/Sc2": 29, "Science Support": 3},
    10: {"10S/Sc1": 29, "10S/Sc2": 28, "10S/Sc3": 23, "10S/Sc4": 22, "10S/Sc5": 12, "Science Support": 3},
    11: {"11S/Sc1": 28, "11S/Sc2": 26, "11S/Sc3": 21, "11S/Sc4": 13, "11S/Sc5": 13, "Science Support": 5},
}
SET_EFFECT = {"Sc1": 0.9, "Sc2": 0.35, "Sc3": -0.1, "Sc4": -0.5, "Sc5": -0.9}
DEMO_TESTS = {
    7: [("7B01", "Cells", 30), ("7C01", "Particles", 25), ("7P01", "Forces", 30), ("7B02", "Organ Systems", 28)],
    8: [("8C01", "Periodic Table", 30), ("8P01", "Energy", 25), ("8B01", "Digestion", 30),
        ("8C02", "Metals and Acids", 26), ("8P02", "Light", 24)],
    9: [("9B01", "Genetics", 30), ("9C01", "Chemical Reactions", 32), ("9P01", "Electricity", 30),
        ("9B02", "Ecology", 25)],
    10: [("4B01", "Cell Biology", 40), ("4C01", "Atomic Structure", 36), ("4P01", "Energy", 35),
         ("4B02", "Organisation", 40), ("4C02", "Bonding", 38)],
    11: [("4B05", "Homeostasis", 40), ("4C09", "Organic Chemistry", 35), ("4P06", "Waves", 38),
         ("4B06", "Inheritance", 42), ("4C10", "Chemical Analysis", 30), ("4P07", "Magnetism", 34)],
}


DEMO_TEACHERS = ["Ms Rahman", "Mr Owen", "Dr Clarke", "Mrs Singh", "Mr Byrne", "Ms Adeyemi", "Mr Walsh",
                 "Miss Patel", "Mr Hughes", "Mrs Evans", "Ms Novak", "Mr Grant"]


def demo_teachers(seed=7):
    rnd = random.Random(seed)
    return {(y, c): rnd.choice(DEMO_TEACHERS) for y in YEARS for c in DEMO_CLASSES[y]}


def demo_data(academic_year, seed=2026):
    rnd = random.Random(seed)
    students, tests = [], []
    start = int(academic_year[:4])
    for y in YEARS:
        roster = []
        for cls, n in DEMO_CLASSES[y].items():
            set_effect = SET_EFFECT.get(cls.split("/")[-1], -1.1)
            for _ in range(n):
                ability = rnd.gauss(set_effect, 0.65)
                maths = None if rnd.random() < 0.08 else int(min(120, max(80, rnd.gauss(104 + 6.5 * ability, 3.2))))
                reading = None if rnd.random() < 0.08 else int(min(120, max(80, rnd.gauss(104 + 6 * ability, 3.8))))
                s = Student(
                    upn=f"X90{rnd.randrange(10**10):010d}", last=rnd.choice(LAST_NAMES),
                    first=rnd.choice(FIRST_NAMES), year=y, sex=rnd.choice("MF"),
                    sen=rnd.choices(["N", "", "K", "E"], [45, 38, 15, 2])[0],
                    pp=rnd.choices(["Y", "N"], [33, 67])[0],
                    att=round(min(100.0, max(55.0, rnd.gauss(94, 6))), 1),
                    ks2=ks2_average(maths, reading), cls=cls)
                s._ability = ability + rnd.gauss(0, 0.25) - (0.15 if s.sen in ("K", "E") else 0)
                roster.append(s)
        students += roster
        classes = list(DEMO_CLASSES[y])
        for n, (code, title, mm) in enumerate(DEMO_TESTS[y]):
            date = dt.date(start, 9, 15) + dt.timedelta(days=14 * n + y)
            difficulty = rnd.uniform(-0.5, 0.6)
            last_test = n == len(DEMO_TESTS[y]) - 1
            unmarked = set(rnd.sample(classes, 2)) if last_test and y in (8, 10) else set()
            t = Test(y, code, title, mm, date)
            for s in roster:
                if s.cls in unmarked:
                    continue
                if rnd.random() < 0.04:
                    t.scores[s.key] = "A"
                    continue
                p = 1 / (1 + math.exp(-(1.1 * s._ability + difficulty + rnd.gauss(0, 0.35))))
                t.scores[s.key] = int(round(mm * min(1, max(0, p))))
            tests.append(t)
    for s in students:
        del s._ability
    return students, tests


# ------------------------------------------------------------------ workbook

class Builder:
    def __init__(self, data: Dataset):
        self.d = data
        self.wb = Workbook()
        base = Font(name=FONT, sz=10, family=2)
        self.wb._fonts = IndexedList([base])
        self.wb._named_styles["Normal"].font = base
        self.wb.code_name = "ThisWorkbook"
        props = self.wb.properties
        props.title = "Science Mastery Quiz Tracker"
        props.subject = f"Mastery quizzes, Years 7 to 11, {data.academic_year}"
        props.creator = props.lastModifiedBy = "Mastery Quiz Tracker generator"
        props.keywords = "stanine; mastery quiz; KS3; KS4; science"
        self.extras = xlsm_package.Extras()
        self.students = {y: sorted((s for s in data.students if s.year == y),
                                   key=lambda s: (s.last.upper(), s.first.upper())) for y in YEARS}
        self.tests = {y: [t for t in data.tests if t.year == y] for y in YEARS}
        self.classes = {y: sorted({s.cls for s in self.students[y] if s.cls}, key=str.upper) for y in YEARS}

    # -- helpers -------------------------------------------------------------
    def band_buttons(self, ws, last_col, buttons, width=100):
        """Buttons right-aligned inside the navy title band (row 1, columns A..last_col)."""
        right = 0
        for c in range(1, last_col + 1):
            dim = ws.column_dimensions.get(get_column_letter(c))
            right += xlsm_package.col_px(dim.width if dim is not None and dim.width else 8.43)
        total = len(buttons) * width + (len(buttons) - 1) * 8
        self.extras.add_buttons(ws, buttons, left_px=right - total - 10, top_px=8, width_px=width,
                                height_px=24)

    def name(self, name, text, sheet=None):
        dn = DefinedName(name, attr_text=text)
        if sheet is None:
            self.wb.defined_names[name] = dn
        else:
            sheet.defined_names[name] = dn

    def build(self):
        wb = self.wb
        self.ws_start = wb.active
        self.ws_start.title = "Start"
        self.ws_over = wb.create_sheet("Overview")
        self.ws_dash = wb.create_sheet("Dashboard")
        self.ws_watch = wb.create_sheet("Watch List")
        self.ws_year = {y: wb.create_sheet(YEAR_NAME[y]) for y in YEARS}
        self.ws_info = wb.create_sheet("Assessment Info")
        self.ws_set = wb.create_sheet("Settings")
        self.ws_calc = wb.create_sheet("Calc")
        codes = {self.ws_start: "shStart", self.ws_over: "shOverview", self.ws_dash: "shDashboard",
                 self.ws_watch: "shWatch", self.ws_info: "shInfo", self.ws_set: "shSettings",
                 self.ws_calc: "shCalc"}
        codes.update({self.ws_year[y]: CODENAME[y] for y in YEARS})
        for ws, code in codes.items():
            ws.sheet_properties.codeName = code
        tabs = {self.ws_start: NAVY, self.ws_over: BLUE, self.ws_dash: BLUE, self.ws_watch: BLUE,
                self.ws_info: MUTED, self.ws_set: MUTED}
        for ws, colour in tabs.items():
            ws.sheet_properties.tabColor = colour

        self.build_calc_lists()
        for y in YEARS:
            self.build_year(y)
        self.build_register()
        self.build_settings()
        self.build_calc_helpers()
        self.build_overview()
        self.build_dashboard()
        self.build_watch()
        self.build_start()
        self.ws_calc.sheet_state = "hidden"
        for ws in (self.ws_start, self.ws_over, self.ws_dash, self.ws_watch, self.ws_info, self.ws_set):
            print_setup(ws, fit_width=True)
        for y in YEARS:
            print_setup(self.ws_year[y], fit_width=False)
            self.ws_year[y].print_title_rows = "1:1"
            self.ws_year[y].print_title_cols = "A:C"
        wb.calculation.fullCalcOnLoad = True
        wb.active = 0
        return wb

    # -- year sheets ---------------------------------------------------------
    def build_year(self, y):
        ws = self.ws_year[y]
        t = TABLE[y]
        students = self.students[y]
        tests = self.tests[y]
        n = len(students)
        last_row = n + 1
        headers = [h for h, _, _ in FIXED]
        for test in tests:
            headers += [raw_header(test.name), stanine_header(test.name)]
        ncols = len(headers)

        ws.row_dimensions[1].height = 168
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(1, c, h)
            cell.border = BOX
            if c <= 3:
                cell.font = font(9, True, WHITE)
                cell.fill = fill(NAVY)
                cell.alignment = Alignment(horizontal="left", vertical="bottom", wrap_text=True)
            elif c <= len(FIXED):
                summary = h in ("Tests Sat", "Mean Stanine", "vs KS2 Band")
                cell.font = font(9, True, WHITE)
                cell.fill = fill(BLUE_500 if summary else NAVY)
                cell.alignment = Alignment(horizontal="center", vertical="bottom",
                                           text_rotation=90, wrap_text=True)
            else:
                k = (c - FIRST_TEST_COL) // 2
                cell.font = font(9, True, INK)
                cell.fill = fill(TEST_FILLS[k % 2])
                cell.alignment = Alignment(horizontal="center", vertical="bottom",
                                           text_rotation=90, wrap_text=True)
        for c, (_, w, _) in enumerate(FIXED, start=1):
            ws.column_dimensions[get_column_letter(c)].width = w
        for c in range(FIRST_TEST_COL, ncols + 1):
            ws.column_dimensions[get_column_letter(c)].width = VBA["TEST_COL_WIDTH"]

        body_font = font(10)
        grey = fill(STANINE_FILL)
        if not students:                        # a table needs one row, even if it is empty
            students = [Student("", "", "", y)]
            n, last_row = 1, 2
        for i, s in enumerate(students):
            r = i + 2
            values = [s.upn or None, s.last or None, s.first or None, s.sex or None, s.sen or None,
                      s.pp or None, s.att, s.ks2 if s.last else None, "=" + ks2_band_formula(t), s.cls or None,
                      "=" + tests_sat_formula(t, r), "=" + mean_stanine_formula(r), "=" + vs_ks2_formula(t)]
            for c, v in enumerate(values, start=1):
                cell = ws.cell(r, c, v)
                cell.font = body_font
                cell.number_format = FIXED[c - 1][2]
                if c >= 4:
                    cell.alignment = CENTER
                if FIXED[c - 1][0] in FORMULA_COLS:
                    cell.fill = grey
            for k, test in enumerate(tests):
                rc = FIRST_TEST_COL + 2 * k
                v = test.scores.get(s.key)
                raw = ws.cell(r, rc, v)
                raw.font = body_font
                raw.alignment = CENTER
                st = ws.cell(r, rc + 1, "=" + stanine_formula(t, raw_header(test.name)))
                st.font = body_font
                st.alignment = CENTER
                st.number_format = "0"
                st.fill = grey

        ref = f"A1:{get_column_letter(ncols)}{last_row}"
        tab = Table(displayName=t, ref=ref)
        cols = []
        for c, h in enumerate(headers, start=1):
            tc = TableColumn(id=c, name=h)
            if h == "Avg KS2 Band":
                tc.calculatedColumnFormula = TableFormula(attr_text=ks2_band_formula(t))
            elif h == "Tests Sat":
                tc.calculatedColumnFormula = TableFormula(attr_text=tests_sat_formula(t, 2))
            elif h == "Mean Stanine":
                tc.calculatedColumnFormula = TableFormula(attr_text=mean_stanine_formula(2))
            elif h == "vs KS2 Band":
                tc.calculatedColumnFormula = TableFormula(attr_text=vs_ks2_formula(t))
            elif h.endswith("\n" + STANINE_SUFFIX):
                raw_name = h[: -len(STANINE_SUFFIX)] + RAW_SUFFIX
                tc.calculatedColumnFormula = TableFormula(attr_text=stanine_formula(t, raw_name))
            cols.append(tc)
        tab.tableColumns = cols
        tab._initialise_columns = lambda: None
        from openpyxl.worksheet.filters import AutoFilter
        tab.autoFilter = AutoFilter(ref=ref)
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
        ws.add_table(tab)
        ws.freeze_panes = "D2"

        # Conditional formats: stanine icons (number + icon), arrows for progress.
        icon_col = {"Avg KS2 Band": (4, 7), "Mean Stanine": (3.5, 6.5)}
        for h, (lo_, hi_) in icon_col.items():
            col = get_column_letter(headers.index(h) + 1)
            ws.conditional_formatting.add(f"{col}2:{col}{last_row}", traffic_lights(lo_, hi_))
        col = get_column_letter(headers.index("vs KS2 Band") + 1)
        ws.conditional_formatting.add(f"{col}2:{col}{last_row}", arrows())
        for k, test in enumerate(tests):
            raw_col = get_column_letter(FIRST_TEST_COL + 2 * k)
            st_col = get_column_letter(FIRST_TEST_COL + 2 * k + 1)
            ws.conditional_formatting.add(f"{st_col}2:{st_col}{last_row}", traffic_lights(4, 7))
            dv = DataValidation(type="custom", formula1=raw_rule(f"{raw_col}2", test.max_marks),
                                allow_blank=True, showErrorMessage=True, showInputMessage=False,
                                errorTitle="Score out of range",
                                error=f"Type a mark from 0 to {test.max_marks:g}, or A for absent. "
                                      "Leave the cell blank if the student has not sat the test yet.")
            ws.add_data_validation(dv)
            dv.add(f"{raw_col}2:{raw_col}{last_row}")

        # Buttons and slicers live in the frozen top-left corner (A1:C1).
        self.extras.add_year_sheet(ws, t, class_col_id=headers.index("Class") + 1,
                                   band_col_id=headers.index("Avg KS2 Band") + 1,
                                   col_widths=[FIXED[i][1] for i in range(3)],
                                   row_height=168, suffix=f"Y{y}")

    # -- register --------------------------------------------------------------
    REG_COLS = ["Year Group", "Code", "Title", "Test Name", "Max Marks", "Date", "Students", "Sat",
                "Absent", "Not Entered", "Completion", "Year Mean", "Year Mean %", "Year SD", "Check",
                "Year Index", "Raw Col", "Seq Key", "Name Key"]
    REG_HIDDEN = ["Year Index", "Raw Col", "Seq Key", "Name Key"]

    def reg_formulas(self):
        T = REG
        yi, rc = tr(T, "Year Index"), tr(T, "Raw Col")
        data = f"INDEX({choose(yi, '#Data')},0,{rc})"
        return {
            "Students": f'IF({yi}=0,"",COUNTA({choose(yi, "Preferred Last name")}))',
            "Sat": f'IF({rc}=0,"",COUNT({data}))',
            "Absent": f'IF({rc}=0,"",COUNTIF({data},"A"))',
            "Not Entered": f'IF({rc}=0,"",{tr(T, "Students")}-{tr(T, "Sat")}-{tr(T, "Absent")})',
            "Completion": f'IF(OR({rc}=0,N({tr(T, "Students")})=0),"",({tr(T, "Sat")}+{tr(T, "Absent")})/{tr(T, "Students")})',
            "Year Mean": f'IF(N({tr(T, "Sat")})=0,"",AVERAGE({data}))',
            "Year Mean %": f'IF(OR({tr(T, "Year Mean")}="",N({tr(T, "Max Marks")})=0),"",{tr(T, "Year Mean")}/{tr(T, "Max Marks")})',
            "Year SD": f'IF(N({tr(T, "Sat")})=0,"",_xlfn.STDEV.P({data}))',
            "Check": (f'IF({tr(T, "Test Name")}="","",IF({yi}=0,"Unknown year group",'
                      f'IF({rc}=0,"Columns not found on the year sheet","OK")))'),
            "Year Index": f'IFERROR(MATCH({tr(T, "Year Group")},YearNames,0),0)',
            "Raw Col": (f'IF({yi}=0,0,IFERROR(MATCH({tr(T, "Test Name")}&{LF}&"{RAW_SUFFIX}",'
                        f'{choose(yi, "#Headers")},0),0))'),
            "Seq Key": (f'{tr(T, "Year Group")}&"|"&COUNTIF(INDEX({T}[Year Group],1):'
                        f'{tr(T, "Year Group")},{tr(T, "Year Group")})'),
            "Name Key": f'{tr(T, "Year Group")}&"|"&{tr(T, "Test Name")}',
        }

    def build_register(self):
        ws = self.ws_info
        ws.sheet_view.showGridLines = False
        cols = self.REG_COLS
        last_col = len(cols)
        title_block(ws, "Assessment register",
                    "One row per test. The Manage tests button adds, edits and removes rows - "
                    "you only need to type here to correct a code, title, maximum mark or date.", 15)
        top = 4
        header_cells(ws, top, 1, cols, height=32)
        formulas = self.reg_formulas()
        tests = [t for y in YEARS for t in self.tests[y]]
        rows = tests or [None]
        for i, test in enumerate(rows):
            r = top + 1 + i
            for c, h in enumerate(cols, start=1):
                cell = ws.cell(r, c)
                cell.font = font(10)
                cell.border = BOX
                if h in formulas:
                    cell.value = "=" + formulas[h]
                    cell.fill = fill(GREY_FILL)
                elif test is not None:
                    cell.value = {"Year Group": YEAR_NAME[test.year], "Code": test.code or None,
                                  "Title": test.title, "Test Name": test.name,
                                  "Max Marks": test.max_marks, "Date": test.date}[h]
                cell.alignment = LEFT if h in ("Title", "Test Name", "Check") else CENTER
                cell.number_format = {"Date": "dd/mm/yyyy", "Completion": "0%", "Year Mean": "0.0",
                                      "Year Mean %": "0%", "Year SD": "0.0", "Max Marks": "0"}.get(h, "General")
        last_row = top + len(rows)
        ref = f"A{top}:{get_column_letter(last_col)}{last_row}"
        tab = Table(displayName=REG, ref=ref)
        tab.tableColumns = [TableColumn(id=i + 1, name=h, calculatedColumnFormula=(
            TableFormula(attr_text=formulas[h]) if h in formulas else None)) for i, h in enumerate(cols)]
        from openpyxl.worksheet.filters import AutoFilter
        tab.autoFilter = AutoFilter(ref=ref)
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
        ws.add_table(tab)
        widths(ws, {"A": 10, "B": 8, "C": 26, "D": 30, "E": 7, "F": 11, "G": 9, "H": 7, "I": 8,
                    "J": 9, "K": 11, "L": 8, "M": 9, "N": 8, "O": 28})
        for h in self.REG_HIDDEN:
            ws.column_dimensions[get_column_letter(cols.index(h) + 1)].hidden = True
        ws.freeze_panes = f"A{top + 1}"
        comp = get_column_letter(cols.index("Completion") + 1)
        rng = f"{comp}{top + 1}:{comp}{top + 400}"
        for rule in completion_rules(f"{comp}{top + 1}"):
            ws.conditional_formatting.add(rng, rule)
        chk = get_column_letter(cols.index("Check") + 1)
        ws.conditional_formatting.add(f"{chk}{top + 1}:{chk}{top + 400}", FormulaRule(
            formula=[f'AND({chk}{top + 1}<>"",{chk}{top + 1}<>"OK")'],
            font=Font(name=FONT, color="B42318", bold=True)))

        # Stanine lookup, used by every Stanine column through the _Stanine name.
        sc = last_col + 2                                   # one spare column after the table
        c1, c2, c3 = (get_column_letter(sc + k) for k in range(3))
        section(ws, f"{c1}{top - 1}", "Stanine lookup")
        header_cells(ws, top, sc, ["z-score from", "Stanine", "Meaning"], height=32)
        for k, z in enumerate(STANINE_Z):
            r = top + 1 + k
            put(ws, f"{c1}{r}", z, al=CENTER, border=BOX, fmt="0.00")
            put(ws, f"{c2}{r}", k + 1, f=font(10, True), al=CENTER, border=BOX)
            put(ws, f"{c3}{r}", STANINE_WORDS[k], al=LEFT, border=BOX)
        self.name("_Stanine", f"'Assessment Info'!${c1}${top + 1}:${c2}${top + 9}")
        put(ws, f"{c1}{top + 11}",
            "How a stanine is worked out: the score is turned into a z-score, "
            "(score - year-group mean) / year-group standard deviation, for that test, "
            "then placed in one of nine bands, each half a standard deviation wide. "
            "5 is average; about 4% of students get a 1 and 4% get a 9.",
            f=font(9, color=INK2), al=WRAP)
        ws.merge_cells(f"{c1}{top + 11}:{c3}{top + 16}")
        ws.column_dimensions[c1].width = 12
        ws.column_dimensions[c2].width = 9
        ws.column_dimensions[c3].width = 22
        self.band_buttons(ws, 15, [("Manage tests", "ShowManageTests"), ("Start", "GoToStart")])

    # -- settings --------------------------------------------------------------
    def build_settings(self):
        ws = self.ws_set
        title_block(ws, "Settings", "Yellow cells are yours to change. Everything else updates itself.", 8)
        widths(ws, {"A": 2, "B": 58, "C": 16, "D": 22, "E": 14, "F": 2, "G": 2, "H": 2})
        section(ws, "B4", "Workbook")
        put(ws, "B5", "Academic year")
        selector(ws, "C5", self.d.academic_year)
        put(ws, "B6", "School or department (shown in titles)")
        selector(ws, "C6", self.d.school)
        self.name("AcademicYear", "Settings!$C$5")
        self.name("SchoolName", "Settings!$C$6")

        section(ws, "B8", "Watch list rules")
        put(ws, "B9", "Flag a student whose mean stanine is at least this far below their KS2 band")
        selector(ws, "C9", 1.5).number_format = "0.0"
        put(ws, "B10", "... once they have sat at least this many tests")
        selector(ws, "C10", 2)
        put(ws, "B11", "Also flag students with no KS2 data whose mean stanine is this or lower")
        selector(ws, "C11", 3).number_format = "0.0"
        self.name("WlGap", "Settings!$C$9")
        self.name("WlMinTests", "Settings!$C$10")
        self.name("WlLowMean", "Settings!$C$11")

        section(ws, "B13", "Year groups")
        header_cells(ws, 14, 2, ["Year group", "Leaves in", "Sheet"])
        start = int(self.d.academic_year[:4])
        for i, y in enumerate(YEARS):
            r = 15 + i
            put(ws, f"B{r}", YEAR_NAME[y], border=BOX)
            put(ws, f"C{r}", start + 1 + (11 - y), al=CENTER, border=BOX, fmt="0")
            put(ws, f"D{r}", YEAR_NAME[y], border=BOX)

        section(ws, "B21", "Class teachers (optional)")
        put(ws, "B22", "Type a teacher's name or initials next to each class. The dashboards show them. "
                       "Add a row at the bottom of the table for a new class.",
            f=font(9, color=INK2), al=WRAP)
        ws.row_dimensions[22].height = 26
        top = 24
        cols = ["Year Group", "Class", "Teacher", "Key"]
        header_cells(ws, top, 2, cols)
        pairs = [(y, c) for y in YEARS for c in self.classes[y]] or [(7, "")]
        for i, (y, c) in enumerate(pairs):
            r = top + 1 + i
            put(ws, f"B{r}", YEAR_NAME[y], border=BOX)
            put(ws, f"C{r}", c or None, border=BOX)
            cell = put(ws, f"D{r}", self.d.teachers.get((y, c)), fl=INPUT_FILL, border=BOX)
            cell.protection = UNLOCKED
            put(ws, f"E{r}", f'={tr(TEACHERS, "Year Group")}&"|"&{tr(TEACHERS, "Class")}',
                f=font(9, color=MUTED), border=BOX)
        last = top + len(pairs)
        ref = f"B{top}:E{last}"
        tab = Table(displayName=TEACHERS, ref=ref)
        tab.tableColumns = [TableColumn(id=i + 1, name=h) for i, h in enumerate(cols)]
        tab.tableColumns[3].calculatedColumnFormula = TableFormula(
            attr_text=f'{tr(TEACHERS, "Year Group")}&"|"&{tr(TEACHERS, "Class")}')
        from openpyxl.worksheet.filters import AutoFilter
        tab.autoFilter = AutoFilter(ref=ref)
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
        ws.add_table(tab)
        ws.column_dimensions["E"].hidden = True
        self.band_buttons(ws, 8, [("Start", "GoToStart")])

    # -- hidden helper sheet ------------------------------------------------------
    def build_calc_lists(self):
        ws = self.ws_calc
        put(ws, "A1", "Helper calculations for the dashboards and drop-down lists. Please do not edit.",
            f=font(10, True, "B42318"))
        for i, y in enumerate(YEARS):
            put(ws, f"B{3 + i}", YEAR_NAME[y])
            put(ws, f"C{3 + i}", TABLE[y])
        self.name("YearNames", "Calc!$B$3:$B$7")
        for i, m in enumerate(MEASURES):
            put(ws, f"E{3 + i}", m)
        self.name("MeasureList", "Calc!$E$3:$E$5")

    def build_calc_helpers(self):
        ws = self.ws_calc
        first, last = 10, 10 + MAX_STUDENTS - 1
        # Distinct, sorted class names per year: mirror the Class column, give each
        # first occurrence its alphabetical rank, then list by rank.
        for i, y in enumerate(YEARS):
            m, k = get_column_letter(8 + 2 * i), get_column_letter(9 + 2 * i)      # H/I ... P/Q
            put(ws, f"{m}{first - 1}", f"{YEAR_NAME[y]} class", f=font(9, True))
            for r in range(first, last + 1):
                ws[f"{m}{r}"] = f'=IFERROR(INDEX({TABLE[y]}[Class],ROWS({m}${first}:{m}{r}))&"","")'
                ws[f"{k}{r}"] = (f'=IF({m}{r}="","",IF(COUNTIF({m}${first}:{m}{r},{m}{r})=1,'
                                 f'COUNTIF({m}${first}:{m}${last},"<"&{m}{r})+1,""))')
            slot = get_column_letter(19 + i)                                        # S ... W
            put(ws, f"{slot}2", YEAR_NAME[y], f=font(9, True))
            for j in range(CLASS_SLOTS):
                ws[f"{slot}{3 + j}"] = (f'=IFERROR(INDEX({m}${first}:{m}${last},MATCH(SMALL({k}${first}:'
                                        f'{k}${last},{j + 1}),{k}${first}:{k}${last},0)),"")')
            ws[f"{slot}12"] = f'=COUNT({k}${first}:{k}${last})'
            self.name(f"ClassSlotsY{y}", f"Calc!${slot}$3:${slot}${2 + CLASS_SLOTS}")

        # Dashboard: selected year group and test.
        self.name("SelYear", "Dashboard!$C$4")
        self.name("SelMeasure", "Dashboard!$C$5")
        self.name("SelTest", "Dashboard!$C$6")
        self.name("SelClass", "Dashboard!$C$7")
        self.name("dYear", "MATCH(SelYear,YearNames,0)")
        cols = {"dData": "#Data", "dHdr": "#Headers", "dClass": "Class", "dLast": "Preferred Last name",
                "dFirst": "Preferred First name", "dMean": "Mean Stanine", "dVs": "vs KS2 Band",
                "dKS2": "Avg KS2 Band", "dPP": "PP Deprivation", "dSEN": "SEN Status Code",
                "dSex": "Sex Code", "dTests": "Tests Sat"}
        for nm, part in cols.items():
            self.name(nm, choose("dYear", part))
        self.name("dRaw", f'INDEX(dData,0,MATCH(SelTest&{LF}&"{RAW_SUFFIX}",dHdr,0))')
        self.name("dStn", f'INDEX(dData,0,MATCH(SelTest&{LF}&"{STANINE_SUFFIX}",dHdr,0))')
        self.name("dHasTest", f'ISNUMBER(MATCH(SelTest&{LF}&"{RAW_SUFFIX}",dHdr,0))')
        self.name("dMax", f'INDEX({REG}[Max Marks],MATCH(SelYear&"|"&SelTest,{REG}[Name Key],0))')

        # Class slots for the selected year group (Y3:Y10) and its test list (AA3:AC62).
        put(ws, "Y2", "Selected year classes", f=font(9, True))
        for j in range(CLASS_SLOTS):
            ws[f"Y{3 + j}"] = f'=IFERROR(INDEX($S{3 + j}:$W{3 + j},dYear),"")'
        ws["Y12"] = '=IFERROR(INDEX($S$12:$W$12,dYear),0)'
        put(ws, "AA2", "Selected year tests", f=font(9, True))
        put(ws, "AB2", "Max", f=font(9, True))
        put(ws, "AC2", "Date", f=font(9, True))
        put(ws, "AD2", "Code", f=font(9, True))
        for j in range(MAX_TESTS):
            r = 3 + j
            key = f'MATCH(SelYear&"|"&{j + 1},{REG}[Seq Key],0)'
            ws[f"AA{r}"] = f'=IFERROR(INDEX({REG}[Test Name],{key})&"","")'
            ws[f"AB{r}"] = f'=IF(AA{r}="","",IFERROR(INDEX({REG}[Max Marks],{key}),""))'
            ws[f"AC{r}"] = (f'=IF(AA{r}="","",IFERROR(IF(INDEX({REG}[Date],{key})="","",'
                            f'INDEX({REG}[Date],{key})),""))')
            ws[f"AD{r}"] = f'=IF(AA{r}="","",IFERROR(INDEX({REG}[Code],{key})&"",""))'
            ws[f"AC{r}"].number_format = "dd/mm/yyyy"
        ws["AA64"] = f'=COUNTIF({REG}[Year Group],SelYear)'
        self.name("dTestCount", "Calc!$AA$64")
        self.name("DashTests", "OFFSET(Calc!$AA$3,0,0,MAX(1,MIN(60,Calc!$AA$64)),1)")
        self.name("DashClasses", "OFFSET(Calc!$Y$3,0,0,MAX(1,Calc!$Y$12),1)")

        # Latest test per year group (Overview, Watch List, Start).
        put(ws, "AF2", "Latest test", f=font(9, True))
        for i, y in enumerate(YEARS):
            r = 3 + i
            cnt = f'COUNTIF({REG}[Year Group],"{YEAR_NAME[y]}")'
            key = f'MATCH("{YEAR_NAME[y]}|"&{cnt},{REG}[Seq Key],0)'
            ws[f"AF{r}"] = f'=IF({cnt}=0,"",IFERROR(INDEX({REG}[Test Name],{key}),""))'
            ws[f"AG{r}"] = f'=IF(AF{r}="","",IFERROR(INDEX({REG}[Max Marks],{key}),""))'
            ws[f"AH{r}"] = (f'=IF(AF{r}="","",IFERROR(IF(INDEX({REG}[Date],{key})="","",'
                            f'INDEX({REG}[Date],{key})),""))')
            ws[f"AH{r}"].number_format = "dd/mm/yyyy"
            ws[f"AI{r}"] = f'=COUNTIF({REG}[Year Group],"{YEAR_NAME[y]}")'

        # Chart feeds. The trend plots mean stanine per test (whole year and focus class).
        put(ws, "AK2", "Trend label", f=font(9, True))
        put(ws, "AL2", "Whole year", f=font(9, True))
        put(ws, "AM2", "Focus class", f=font(9, True))
        for j in range(MAX_TESTS):
            r = 3 + j
            stn = f'INDEX(dData,0,MATCH(AA{r}&{LF}&"{STANINE_SUFFIX}",dHdr,0))'
            ws[f"AK{r}"] = f'=IF(AA{r}="","",IF(AD{r}<>"",AD{r},LEFT(AA{r},14)))'
            ws[f"AL{r}"] = f'=IF(AA{r}="",NA(),IFERROR(AVERAGE({stn}),NA()))'
            ws[f"AM{r}"] = f'=IF(OR(AA{r}="",SelClass=""),NA(),IFERROR(AVERAGEIFS({stn},dClass,SelClass),NA()))'
        n_tests = "MAX(1,MIN(60,Calc!$AA$64))"
        self.name("chTrendCats", f"OFFSET(Calc!$AK$3,0,0,{n_tests},1)", sheet=ws)
        self.name("chTrendYear", f"OFFSET(Calc!$AL$3,0,0,{n_tests},1)", sheet=ws)
        self.name("chTrendClass", f"OFFSET(Calc!$AM$3,0,0,{n_tests},1)", sheet=ws)

        put(ws, "AO2", "Test by class", f=font(9, True))
        ws["AO3"] = "Whole year"
        ws["AP3"] = '=IFERROR(AVERAGE(dRaw)/dMax,NA())'
        for r in range(3, 4 + CLASS_SLOTS):
            ws[f"AP{r}"].number_format = "0%"
            ws[f"AS{r}"].number_format = "0%"
            ws[f"AT{r}"].number_format = "0%"
        for j in range(CLASS_SLOTS):
            r = 4 + j
            ws[f"AO{r}"] = f'=Y{3 + j}'
            ws[f"AP{r}"] = f'=IF(AO{r}="",NA(),IFERROR(AVERAGEIFS(dRaw,dClass,AO{r})/dMax,NA()))'
        self.name("chClassCats", "OFFSET(Calc!$AO$3,0,0,1+Calc!$Y$12,1)", sheet=ws)
        self.name("chClassVals", "OFFSET(Calc!$AP$3,0,0,1+Calc!$Y$12,1)", sheet=ws)

        put(ws, "AR2", "Stanine", f=font(9, True))
        put(ws, "AS2", "This test", f=font(9, True))
        put(ws, "AT2", "Expected", f=font(9, True))
        for k in range(9):
            r = 3 + k
            ws[f"AR{r}"] = k + 1
            ws[f"AS{r}"] = f'=IFERROR(COUNTIF(dStn,AR{r})/COUNT(dStn),0)'
            ws[f"AT{r}"] = EXPECTED_PCT[k] / 100

        # Watch list: mirror the chosen year group, flag, then sort by how far below KS2.
        self.name("WlYear", "'Watch List'!$C$4")
        self.name("wYear", "MATCH(WlYear,YearNames,0)")
        wcols = {"wData": "#Data", "wHdr": "#Headers", "wLast": "Preferred Last name",
                 "wFirst": "Preferred First name", "wClass": "Class", "wKS2": "Avg KS2 Band",
                 "wTests": "Tests Sat", "wMean": "Mean Stanine", "wVs": "vs KS2 Band"}
        for nm, part in wcols.items():
            self.name(nm, choose("wYear", part))
        self.name("wLatest", "INDEX(Calc!$AF$3:$AF$7,wYear)")
        heads = ["Last", "First", "Class", "KS2 band", "Tests", "Mean", "vs KS2", "Latest", "Reason", "Key"]
        for c, h in enumerate(heads):
            put(ws, f"{get_column_letter(53 + c)}69", h, f=font(9, True))            # BA ...
        wfirst, wlast = 70, 70 + MAX_STUDENTS - 1
        for r in range(wfirst, wlast + 1):
            i = r - wfirst + 1
            ws[f"BA{r}"] = f'=IFERROR(INDEX(wLast,{i})&"","")'
            ws[f"BB{r}"] = f'=IF(BA{r}="","",IFERROR(INDEX(wFirst,{i})&"",""))'
            ws[f"BC{r}"] = f'=IF(BA{r}="","",IFERROR(INDEX(wClass,{i})&"",""))'
            ws[f"BD{r}"] = f'=IF(BA{r}="","",IFERROR(INDEX(wKS2,{i})+0,""))'
            ws[f"BE{r}"] = f'=IF(BA{r}="","",IFERROR(INDEX(wTests,{i})+0,0))'
            ws[f"BF{r}"] = f'=IF(BA{r}="","",IFERROR(INDEX(wMean,{i})+0,""))'
            ws[f"BG{r}"] = f'=IF(BA{r}="","",IFERROR(INDEX(wVs,{i})+0,""))'
            ws[f"BH{r}"] = (f'=IF(OR(BA{r}="",wLatest=""),"",IFERROR(INDEX(wData,{i},MATCH(wLatest&{LF}&'
                            f'"{STANINE_SUFFIX}",wHdr,0))+0,""))')
            ws[f"BI{r}"] = (f'=IF(BA{r}="","",IF(AND(ISNUMBER(BG{r}),N(BE{r})>=WlMinTests,BG{r}<=-WlGap),'
                            f'"Below KS2 starting point",IF(AND(NOT(ISNUMBER(BD{r})),ISNUMBER(BF{r}),'
                            f'N(BE{r})>=WlMinTests,BF{r}<=WlLowMean),"Low results, no KS2 data","")))')
            ws[f"BJ{r}"] = (f'=IF(BI{r}="","",IF(BI{r}="Below KS2 starting point",BG{r},100+BF{r})'
                            f'+ROW()/1000000)')
        ws["BA67"] = f'=COUNT(BJ{wfirst}:BJ{wlast})'
        self.name("wlCount", "Calc!$BA$67")
        for k in range(WATCH_ROWS):
            r = 3 + k
            ws[f"BL{r}"] = (f'=IFERROR(MATCH(SMALL($BJ${wfirst}:$BJ${wlast},{k + 1}),'
                            f'$BJ${wfirst}:$BJ${wlast},0),"")')

    # -- dashboards ------------------------------------------------------------
    def teacher_lookup(self, year_expr, class_ref):
        return (f'IF({class_ref}="","",IFERROR(INDEX({TEACHERS}[Teacher],MATCH({year_expr}&"|"&{class_ref},'
                f'{TEACHERS}[Key],0))&"",""))')

    def build_overview(self):
        ws = self.ws_over
        title_block(ws, "Class overview",
                    "How every class in every year group is doing, from all the tests entered so far. "
                    "Blue = above the year group, red = below.", 14)
        widths(ws, {"A": 2, "B": 20, "C": 14, "D": 9, "E": 9, "F": 9, "G": 9, "H": 9, "I": 10,
                    "J": 10, "K": 10, "L": 10, "M": 11, "N": 2})
        tiles = [("Students", "=" + "+".join(f"COUNTA({TABLE[y]}[Preferred Last name])" for y in YEARS), "0"),
                 ("Tests set", f"=COUNTA({REG}[Test Name])", "0"),
                 ("Scores entered", f"=SUM({REG}[Sat])", "#,##0"),
                 ("Tests still being marked", f'=COUNTIFS({REG}[Completion],"<1",{REG}[Test Name],"<>")', "0")]
        stat_tiles(ws, 4, ["B", "D", "G", "J"], ["C", "F", "I", "M"], tiles)

        heads = ["Class", "Teacher", "Students", "KS2 band\n(mean)", "Tests sat\n(mean)", "Mean\nstanine",
                 "vs KS2\nband", "Results at\nstanine 7-9", "Results at\nstanine 1-3",
                 "Latest test\nmean %", "Latest test\nmean stanine", "Latest test\ncompletion"]
        row = 9
        for i, y in enumerate(YEARS):
            t = TABLE[y]
            yname = YEAR_NAME[y]
            latest, latest_max = f"Calc!$AF${3 + i}", f"Calc!$AG${3 + i}"
            band = row
            ws.row_dimensions[band].height = 22
            for col in range(2, 14):
                ws.cell(band, col).fill = fill(BLUE_100)
            put(ws, f"B{band}", yname, f=font(12, True, NAVY), fl=BLUE_100)
            put(ws, f"C{band}", f'="Leaves "&Settings!$C${15 + i}&"   ·   "&COUNTA({t}[Preferred Last name])'
                                f'&" students   ·   "&Calc!$AI${3 + i}&" tests"',
                f=font(9, False, INK2), fl=BLUE_100)
            put(ws, f"H{band}", f'=IF({latest}="","No tests yet","Latest test: "&{latest}'
                                f'&IF(ISNUMBER(Calc!$AH${3 + i}),"  ("&TEXT(Calc!$AH${3 + i},"dd/mm/yyyy")&")",""))',
                f=font(9, True, INK2), fl=BLUE_100)
            header_cells(ws, band + 1, 2, heads, height=30)
            stn_mask = f'(RIGHT({t}[#Headers],8)={LF}&"{STANINE_SUFFIX}")'
            raw_col = f'INDEX({t}[#Data],0,MATCH({latest}&{LF}&"{RAW_SUFFIX}",{t}[#Headers],0))'
            stn_col = f'INDEX({t}[#Data],0,MATCH({latest}&{LF}&"{STANINE_SUFFIX}",{t}[#Headers],0))'
            year_row = band + 2
            first_class_row = year_row + 1
            last_class_row = first_class_row + CLASS_SLOTS - 1
            for j in range(CLASS_SLOTS + 1):
                r = year_row + j
                is_year = j == 0
                if is_year:
                    put(ws, f"B{r}", "Whole year", f=font(10, True), border=YEAR_RULE)
                    crit, mask, count = "", "", f"COUNTA({t}[Preferred Last name])"
                else:
                    put(ws, f"B{r}", f"=Calc!{get_column_letter(19 + i)}{2 + j}")
                    crit, mask, count = f",{t}[Class],$B{r}", f"({t}[Class]=$B{r})*", f"COUNTIF({t}[Class],$B{r})"
                avg = "AVERAGEIFS" if crit else "AVERAGE"

                def share(cond):
                    return (f"SUMPRODUCT({mask}{stn_mask}*ISNUMBER({t}[#Data])*({t}[#Data]{cond}))"
                            f"/SUMPRODUCT({mask}{stn_mask}*ISNUMBER({t}[#Data]))")
                cells = {
                    "C": "" if is_year else self.teacher_lookup(f'"{yname}"', f"$B{r}"),
                    "D": count,
                    "E": f"{avg}({t}[Avg KS2 Band]{crit})",
                    "F": f"{avg}({t}[Tests Sat]{crit})",
                    "G": f"{avg}({t}[Mean Stanine]{crit})",
                    "H": f"{avg}({t}[vs KS2 Band]{crit})",
                    "I": share(">=7"),
                    "J": share("<=3"),
                    "K": f'IF({latest}="","",{avg}({raw_col}{crit})/{latest_max})',
                    "L": f'IF({latest}="","",{avg}({stn_col}{crit}))',
                    "M": f'IF({latest}="","",(COUNTIFS({raw_col},">=0"{crit})+COUNTIFS({raw_col},"A"{crit}))/{count})',
                }
                fmts = {"D": "0", "E": "0.0", "F": "0.0", "G": "0.0", "H": "+0.0;-0.0;0.0", "I": "0%",
                        "J": "0%", "K": "0%", "L": "0.0", "M": "0%"}
                for col, fml in cells.items():
                    if not fml:
                        continue
                    if col != "C":
                        fml = f'IFERROR({fml},"")'
                        if not is_year:
                            fml = f'IF($B{r}="","",{fml})'
                    c = put(ws, f"{col}{r}", "=" + fml, f=font(10, is_year), al=LEFT if col == "C" else CENTER,
                            fmt=fmts.get(col))
                    if is_year:
                        c.border = TOP_RULE
            cls_rows = f"{first_class_row}:{last_class_row}"
            add_diverging(ws, f"G{first_class_row}:G{last_class_row}", f"G{first_class_row}", "5", 0.4, 1.0)
            add_diverging(ws, f"H{first_class_row}:H{last_class_row}", f"H{first_class_row}", "0", 0.2, 0.5)
            add_diverging(ws, f"K{first_class_row}:K{last_class_row}", f"K{first_class_row}", f"$K${year_row}",
                          0.03, 0.10)
            ws.conditional_formatting.add(f"I{year_row}:I{last_class_row}", data_bar(BLUE_200))
            ws.conditional_formatting.add(f"J{year_row}:J{last_class_row}", data_bar(RED_200))
            for rule in completion_rules(f"M{year_row}"):
                ws.conditional_formatting.add(f"M{year_row}:M{last_class_row}", rule)
            row = last_class_row + 2
        put(ws, f"B{row}", "How to read this: 'Mean stanine' compares a class with its whole year group "
                           "(5 = the year average). 'vs KS2 band' compares the class's stanines with the "
                           "same students' KS2 bands, so it shows progress from their starting point. "
                           "Results at 7-9 and 1-3 count every test result, not students.",
            f=font(9, color=INK2), al=WRAP)
        ws.merge_cells(f"B{row}:M{row + 3}")
        ws.freeze_panes = "A8"
        self.band_buttons(ws, 14, [("Dashboard", "GoToDashboard"), ("Watch list", "GoToWatchList"),
                                   ("Start", "GoToStart")])
        protect(ws)

    def build_dashboard(self):
        ws = self.ws_dash
        title_block(ws, "Assessment dashboard",
                    "Choose a year group and a test. Colours compare each class with the whole year group: "
                    "blue = above, red = below.", 26)
        widths(ws, {"A": 2, "B": 30, "C": 11, "D": 9, "E": 9, "F": 10, "G": 10, "H": 10, "I": 10,
                    "J": 10, "K": 10, "L": 10, "M": 10, "N": 10, "O": 10, "P": 2})
        for col in range(17, 28):
            ws.column_dimensions[get_column_letter(col)].width = 9

        y0 = self.d.dashboard_year
        tests0 = self.tests[y0]
        classes0 = self.classes[y0]
        put(ws, "B4", "Year group", f=font(10, True))
        selector(ws, "C4", YEAR_NAME[y0], "D4")
        put(ws, "B5", "Measure in the grid below", f=font(10, True))
        selector(ws, "C5", self.d.dashboard_measure, "D5")
        put(ws, "B6", "Test", f=font(10, True))
        test0 = self.d.dashboard_test if self.d.dashboard_test is not None else (tests0[-1].name if tests0 else "")
        selector(ws, "C6", test0, "G6")
        put(ws, "B7", "Focus class (trend chart)", f=font(10, True))
        class0 = self.d.dashboard_class if self.d.dashboard_class is not None else (classes0[0] if classes0 else "")
        selector(ws, "C7", class0, "E7")
        dv_list(ws, "C4", "=YearNames")
        dv_list(ws, "C5", "=MeasureList")
        dv_list(ws, "C6", "=DashTests")
        dv_list(ws, "C7", "=DashClasses")

        tiles = [("Students", "=IFERROR(COUNTA(dLast),0)", "0"),
                 ("Tests set", "=IFERROR(dTestCount,0)", "0"),
                 ("Completion, this test", "=IFERROR((COUNT(dRaw)+COUNTIF(dRaw,\"A\"))/COUNTA(dLast),\"-\")", "0%"),
                 ("Year mean, this test", "=IFERROR(AVERAGE(dRaw)/dMax,\"-\")", "0%")]
        stat_tiles(ws, 4, ["I", "K", "M", "O"], ["J", "L", "N", "Q"], tiles, height_rows=4)

        # 1. Selected test by class ------------------------------------------------------
        section(ws, "B9", "1.  How each class did on the selected test")
        put(ws, "B10", '=IF(SelTest="","Choose a test above.",SelTest&IFERROR("   ·   out of "&dMax,""))',
            f=font(10, False, INK2, italic=True))
        heads = ["Class", "Teacher", "Students", "Sat", "Absent", "Not\nentered", "Mean\nmark", "Mean %",
                 "Mean\nstanine", "Stanine\n1-3", "Stanine\n4-6", "Stanine\n7-9", "vs year\n(% points)"]
        header_cells(ws, 11, 2, heads, height=30)
        for j in range(CLASS_SLOTS + 1):
            r = 12 + j
            is_year = j == 0
            if is_year:
                put(ws, f"B{r}", "Whole year", f=font(10, True), fl=GREY_FILL, border=YEAR_RULE)
            else:
                put(ws, f"B{r}", f"=Calc!Y{2 + j}")
            b = f"$B{r}"
            skip = f'OR({b}="",N($E{r})=0)' if not is_year else f'N($E{r})=0'

            def count(*crit):
                """COUNTIFS for this class (or the whole year) with extra criteria pairs."""
                pairs = ([] if is_year else ["dClass", b]) + list(crit)
                return ("COUNTIFS(" + ",".join(pairs) + ")") if pairs else "COUNTA(dLast)"

            def mean(rng):
                return f"AVERAGE({rng})" if is_year else f"AVERAGEIFS({rng},dClass,{b})"
            cells = {
                "C": "" if is_year else self.teacher_lookup("SelYear", b),
                "D": "COUNTA(dLast)" if is_year else f'IF({b}="","",COUNTIF(dClass,{b}))',
                "E": 'IF(dHasTest,COUNT(dRaw),"")' if is_year else f'IF({b}="","",{count("dRaw", chr(34) + ">=0" + chr(34))})',
                "F": 'COUNTIF(dRaw,"A")' if is_year else f'IF({b}="","",{count("dRaw", chr(34) + "A" + chr(34))})',
                "G": f'IF($D{r}="","",$D{r}-$E{r}-$F{r})',
                "H": f'IF({skip},"",{mean("dRaw")})',
                "I": f'IF($H{r}="","",$H{r}/dMax)',
                "J": f'IF({skip},"",{mean("dStn")})',
                "K": f'IF({skip},"",{count("dStn", Q("<=3"))}/$E{r})',
                "L": f'IF({skip},"",{count("dStn", Q(">=4"), "dStn", Q("<=6"))}/$E{r})',
                "M": f'IF({skip},"",{count("dStn", Q(">=7"))}/$E{r})',
                "N": "" if is_year else f'IF(OR($I{r}="",$I$12=""),"",($I{r}-$I$12)*100)',
            }
            fmts = {"D": "0", "E": "0", "F": "0", "G": "0", "H": "0.0", "I": "0%", "J": "0.0", "K": "0%",
                    "L": "0%", "M": "0%", "N": "+0;-0;0"}
            for col, fml in cells.items():
                if not fml:
                    continue
                c = put(ws, f"{col}{r}", f'=IFERROR({fml},"")', f=font(10, is_year),
                        al=LEFT if col == "C" else CENTER, fmt=fmts.get(col))
                if is_year:
                    c.border = YEAR_RULE
                    c.fill = fill(GREY_FILL)
            if is_year:
                for col in "CN":
                    ws[f"{col}{r}"].fill = fill(GREY_FILL)
                    ws[f"{col}{r}"].border = YEAR_RULE
        add_diverging(ws, "I13:I20", "I13", "$I$12", 0.03, 0.10)
        add_diverging(ws, "J13:J20", "J13", "$J$12", 0.4, 1.0)
        for rule in completion_rules_counts("G12"):
            ws.conditional_formatting.add("G12:G20", rule)

        # 2. Stanine spread and 3. groups ------------------------------------------------
        section(ws, "B23", "2.  Stanine spread on the selected test")
        header_cells(ws, 24, 2, ["Stanine"] + [str(k) for k in range(1, 10)], height=20)
        put(ws, "B25", "Students", f=font(10, True))
        put(ws, "B26", "% of those who sat", f=font(10, True))
        put(ws, "B27", "Expected % (a normal spread)", f=font(10, False, INK2))
        for k in range(9):
            col = get_column_letter(3 + k)
            put(ws, f"{col}25", f'=IFERROR(COUNTIF(dStn,{k + 1}),"")', al=CENTER, fmt="0")
            put(ws, f"{col}26", f'=IFERROR(COUNTIF(dStn,{k + 1})/COUNT(dStn),"")', al=CENTER, fmt="0%")
            put(ws, f"{col}27", EXPECTED_PCT[k] / 100, f=font(10, color=INK2), al=CENTER, fmt="0%")

        section(ws, "B30", "3.  Groups: the selected test and all tests so far")
        gheads = ["Group", "Students", "Sat this\ntest", "Mean %\n(this test)", "Mean stanine\n(this test)",
                  "Mean stanine\n(all tests)", "vs KS2 band\n(all tests)"]
        header_cells(ws, 31, 2, gheads, height=30)
        groups = [  # (label, range, codes, everyone except those codes?)
            ("Disadvantaged (PP)", "dPP", ["Y"], False), ("Not disadvantaged", "dPP", ["Y"], True), None,
            ("SEN support or EHCP (K, E)", "dSEN", ["K", "E"], False),
            ("No SEN support (N or blank)", "dSEN", ["K", "E"], True), None,
            ("Female", "dSex", ["F"], False), ("Male", "dSex", ["M"], False), None,
        ]
        for i, g in enumerate(groups):
            r = 32 + i
            if g is None:
                a, b = r - 2, r - 1
                label = {34: "Gap: PP minus not PP", 37: "Gap: SEN minus no SEN", 40: "Gap: female minus male"}[r]
                put(ws, f"B{r}", label, f=font(10, True, INK2), border=TOP_RULE)
                gaps = {"E": f'IF(OR($E{a}="",$E{b}=""),"",($E{a}-$E{b})*100)',
                        "F": f'IF(OR($F{a}="",$F{b}=""),"",$F{a}-$F{b})',
                        "G": f'IF(OR($G{a}="",$G{b}=""),"",$G{a}-$G{b})',
                        "H": f'IF(OR($H{a}="",$H{b}=""),"",$H{a}-$H{b})'}
                gfmt = {"E": '+0.0" pts";-0.0" pts";0.0" pts"', "F": "+0.0;-0.0;0.0", "G": "+0.0;-0.0;0.0",
                        "H": "+0.0;-0.0;0.0"}
                for col, fml in gaps.items():
                    put(ws, f"{col}{r}", f'=IFERROR({fml},"")', f=font(10, True, INK2), al=CENTER,
                        fmt=gfmt[col], border=TOP_RULE)
                for col in "CD":
                    ws[f"{col}{r}"].border = TOP_RULE
                continue
            label, rng, codes, others = g

            def total(template, everyone):
                """Sum a SUMIFS/COUNTIFS over the group's codes, or everyone minus those codes.

                "Everyone minus" avoids matching blank cells, which LibreOffice and Excel
                treat differently at the end of a range.
                """
                part = "+".join(template.format(rng=rng, v=Q(v)) for v in codes)
                return f"({everyone}-({part}))" if others else f"({part})"
            students = total("COUNTIFS({rng},{v})", "COUNTA(dLast)")
            sat = total('COUNTIFS({rng},{v},dRaw,">=0")', "COUNT(dRaw)")
            put(ws, f"B{r}", label)
            cells = {
                "C": students,
                "D": sat,
                "E": f'IF({sat}=0,"",{total("SUMIFS(dRaw,{rng},{v})", "SUM(dRaw)")}/{sat}/dMax)',
                "F": (f'IF({sat}=0,"",{total("SUMIFS(dStn,{rng},{v})", "SUM(dStn)")}/'
                      f'{total("COUNTIFS({rng},{v},dStn," + Q(">=1") + ")", "COUNT(dStn)")})'),
                "G": (f'{total("SUMIFS(dMean,{rng},{v})", "SUM(dMean)")}/'
                      f'{total("COUNTIFS({rng},{v},dMean," + Q(">=1") + ")", "COUNT(dMean)")}'),
                "H": (f'{total("SUMIFS(dVs,{rng},{v})", "SUM(dVs)")}/'
                      f'{total("COUNTIFS({rng},{v},dVs," + Q(">=-9") + ")", "COUNT(dVs)")}'),
            }
            fmts = {"C": "0", "D": "0", "E": "0%", "F": "0.0", "G": "0.0", "H": "+0.0;-0.0;0.0"}
            for col, fml in cells.items():
                put(ws, f"{col}{r}", f'=IFERROR({fml},"")', al=CENTER, fmt=fmts[col])

        # 4. Every test: grid ------------------------------------------------------------
        section(ws, "B43", '4.  Every test for the year group')
        put(ws, "B44", '="Showing "&LOWER(SelMeasure)&". Class cells: blue = above the whole year group, '
                       'red = below. Completion: amber = partly entered, red = nothing entered yet."',
            f=font(9, False, INK2, italic=True))
        gheads = ["Test", "Date", "Max", "Sat", "Completion", "Whole year"]
        header_cells(ws, 45, 2, gheads, height=36)
        for j in range(CLASS_SLOTS):
            c = ws.cell(45, 8 + j, f'=IF(Calc!Y{3 + j}="","",Calc!Y{3 + j})')
            c.font = font(9, True, WHITE)
            c.fill = fill(NAVY)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = BOX
        mi = "MATCH(SelMeasure,MeasureList,0)"
        for j in range(MAX_TESTS):
            r = 46 + j
            raw = f'INDEX(dData,0,MATCH($B{r}&{LF}&"{RAW_SUFFIX}",dHdr,0))'
            stn = f'INDEX(dData,0,MATCH($B{r}&{LF}&"{STANINE_SUFFIX}",dHdr,0))'
            put(ws, f"B{r}", f"=Calc!AA{3 + j}")
            put(ws, f"C{r}", f"=Calc!AC{3 + j}", al=CENTER, fmt="dd/mm/yy")
            put(ws, f"D{r}", f"=Calc!AB{3 + j}", al=CENTER, fmt="0")
            put(ws, f"E{r}", f'=IF($B{r}="","",IFERROR(COUNT({raw}),""))', al=CENTER, fmt="0")
            put(ws, f"F{r}", f'=IF($B{r}="","",IFERROR(($E{r}+COUNTIF({raw},"A"))/COUNTA(dLast),""))',
                al=CENTER, fmt="0%")
            put(ws, f"G{r}", f'=IF($B{r}="","",IFERROR(CHOOSE({mi},AVERAGE({raw})/$D{r},AVERAGE({stn}),$F{r}),""))',
                f=font(10, True), al=CENTER, fmt="0%")
            for k in range(CLASS_SLOTS):
                col = get_column_letter(8 + k)
                h = f"{col}$45"
                fml = (f'=IF(OR($B{r}="",{h}=""),"",IFERROR(CHOOSE({mi},AVERAGEIFS({raw},dClass,{h})/$D{r},'
                       f'AVERAGEIFS({stn},dClass,{h}),(COUNTIFS(dClass,{h},{raw},">=0")+COUNTIFS(dClass,{h},{raw},"A"))'
                       f'/COUNTIF(dClass,{h})),""))')
                put(ws, f"{col}{r}", fml, al=CENTER, fmt="0%")
        grid = f"G46:O{45 + MAX_TESTS}"
        ws.conditional_formatting.add("H45:O45", Rule(type="expression", formula=['H45=""'], stopIfTrue=True,
                                                      dxf=DifferentialStyle(fill=PatternFill(bgColor=WHITE, fill_type="solid"))))
        add_grid_rules(ws, f"H46:O{45 + MAX_TESTS}", "H46", "$G46")
        for rule in completion_rules("F46"):
            ws.conditional_formatting.add(f"F46:F{45 + MAX_TESTS}", rule)
        add_number_format_rule(ws, grid, 'SelMeasure="Mean stanine"', "0.0")
        ws.conditional_formatting.add(f"B46:O{45 + MAX_TESTS}", row_rule("$B46"))
        ws.freeze_panes = "A8"

        # Charts
        self.band_buttons(ws, 16, [("Manage tests", "ShowManageTests"), ("Overview", "GoToOverview"),
                                   ("Watch list", "GoToWatchList"), ("Start", "GoToStart")])
        ws.add_chart(class_bar_chart(), "R9")
        ws.add_chart(stanine_chart(), "R23")
        ws.add_chart(trend_chart(), "R43")
        protect(ws)

    def build_watch(self):
        ws = self.ws_watch
        title_block(ws, "Students to check",
                    "Students whose results so far are well below their KS2 starting point. "
                    "Change the rules on the Settings sheet.", 11)
        widths(ws, {"A": 2, "B": 28, "C": 16, "D": 10, "E": 9, "F": 10, "G": 10, "H": 12, "I": 30, "J": 2})
        put(ws, "B4", "Year group", f=font(10, True))
        selector(ws, "C4", YEAR_NAME[self.d.watch_year], "D4")
        dv_list(ws, "C4", "=YearNames")
        put(ws, "B5", '="Rule: mean stanine at least "&TEXT(WlGap,"0.0")&" below the KS2 band after "&WlMinTests'
                      '&" or more tests; or, with no KS2 data, a mean stanine of "&TEXT(WlLowMean,"0.0")&" or lower."',
            f=font(9, False, INK2, italic=True))
        put(ws, "B6", '=IF(wlCount=0,"Nobody is flagged in "&WlYear&" at the moment.",wlCount&" student"'
                      '&IF(wlCount=1,"","s")&" flagged in "&WlYear&IF(wlCount>50,", showing the first 50.","."))',
            f=font(11, True, NAVY))
        heads = ["Student", "Class", "KS2 band", "Tests sat", "Mean\nstanine", "vs KS2\nband",
                 "Latest test\nstanine", "Reason"]
        header_cells(ws, 8, 2, heads, height=30)
        cols = {"C": "BC", "D": "BD", "E": "BE", "F": "BF", "G": "BG", "H": "BH", "I": "BI"}
        fmts = {"D": "0", "E": "0", "F": "0.0", "G": "+0.0;-0.0;0.0", "H": "0"}
        for k in range(WATCH_ROWS):
            r = 9 + k
            idx = f"Calc!$BL${3 + k}"
            put(ws, f"B{r}", f'=IF({idx}="","",INDEX(Calc!$BA$70:$BA$669,{idx})&", "&INDEX(Calc!$BB$70:$BB$669,{idx}))')
            for col, src in cols.items():
                put(ws, f"{col}{r}", f'=IF({idx}="","",INDEX(Calc!${src}$70:${src}$669,{idx}))',
                    al=LEFT if col in "CI" else CENTER, fmt=fmts.get(col))
        ws.conditional_formatting.add(f"G9:G{8 + WATCH_ROWS}", FormulaRule(
            formula=["AND(ISNUMBER(G9),G9<=-1)"], font=Font(name=FONT, color="B42318", bold=True)))
        ws.conditional_formatting.add(f"H9:H{8 + WATCH_ROWS}", traffic_lights(4, 7))
        ws.conditional_formatting.add(f"B9:I{8 + WATCH_ROWS}", row_rule("$B9"))
        ws.freeze_panes = "A9"
        self.band_buttons(ws, 10, [("Dashboard", "GoToDashboard"), ("Overview", "GoToOverview"),
                                   ("Start", "GoToStart")])
        protect(ws)

    def build_start(self):
        ws = self.ws_start
        title_block(ws, "Science Mastery Quiz Tracker", "", 12)
        put(ws, "B2", '=IF(SchoolName="","",SchoolName&"   ·   ")&"Years 7 to 11   ·   "&AcademicYear',
            f=font(11, False, INK2))
        widths(ws, {"A": 2, "B": 22, "C": 12, "D": 12, "E": 12, "F": 14, "G": 34, "H": 16, "I": 2,
                    "J": 12, "K": 12, "L": 2})
        row = 3
        if self.d.demo:
            put(ws, "B3", "DEMO FILE - every student, class and score in this workbook is made up.",
                f=font(11, True, WHITE), fl="B42318")
            for col in "CDEFGH":
                ws[f"{col}3"].fill = fill("B42318")
        # buttons: row 4-5
        self.extras.add_buttons(ws, [("Manage tests", "ShowManageTests"),
                                     ("Update students", "UpdateStudentsFromExport"),
                                     ("Class overview", "GoToOverview"), ("Dashboard", "GoToDashboard"),
                                     ("Watch list", "GoToWatchList"), ("Settings", "GoToSettings")],
                                left_px=18, top_px=96, width_px=128, gap_px=10, height_px=34)
        ws.row_dimensions[4].height = 20
        ws.row_dimensions[5].height = 20
        section(ws, "B7", "Year groups")
        header_cells(ws, 8, 2, ["Year group", "Leaves in", "Students", "Tests set", "Scores entered",
                                "Latest test", "Open"], height=22)
        for i, y in enumerate(YEARS):
            r = 9 + i
            t = TABLE[y]
            put(ws, f"B{r}", YEAR_NAME[y], f=font(10, True), border=BOX)
            put(ws, f"C{r}", f"=Settings!C{15 + i}", al=CENTER, fmt="0", border=BOX)
            put(ws, f"D{r}", f"=COUNTA({t}[Preferred Last name])", al=CENTER, border=BOX)
            put(ws, f"E{r}", f"=Calc!AI{3 + i}", al=CENTER, border=BOX)
            put(ws, f"F{r}", f'=SUMPRODUCT((RIGHT({t}[#Headers],10)={LF}&"{RAW_SUFFIX}")*ISNUMBER({t}[#Data]))',
                al=CENTER, fmt="#,##0", border=BOX)
            put(ws, f"G{r}", f'=IF(Calc!AF{3 + i}="","-",Calc!AF{3 + i})', border=BOX)
            put(ws, f"H{r}", f'=HYPERLINK("#\'{YEAR_NAME[y]}\'!A1","Open {YEAR_NAME[y]}")',
                f=font(10, False, BLUE_500), border=BOX)

        guide = [
            ("For class teachers", [
                "Open your year group's sheet and click your class in the Class slicer (top left).",
                "Type each student's mark in the test's Raw Score column. Type A for absent; leave the cell "
                "empty if they have not sat it yet.",
                "The Stanine column fills itself in. Do not type in the grey columns.",
                "This works in Excel on the web too. Only the buttons need the desktop app.",
            ]),
            ("For the data lead", [
                "Manage tests adds a test to a year group: two new columns plus a row in the register. "
                "The same form edits a title, maximum mark or date, and removes a test added by mistake.",
                "Update students reads a new 'All Students' export from the MIS: new students are added, classes and "
                "details are updated, scores are never touched. Students missing from the export are listed, "
                "not deleted.",
                "Class overview, Dashboard and Watch list update by themselves as marks go in.",
                "Type teachers' names on the Settings sheet to show them on the dashboards.",
            ]),
        ]
        row = 16
        for heading, items in guide:
            section(ws, f"B{row}", heading)
            row += 1
            for n, text in enumerate(items, start=1):
                paragraph(ws, row, f"{n}.   {text}")
                row += 1
            row += 1

        section(ws, f"B{row}", "Reading stanines")
        row += 1
        header_cells(ws, row, 2, ["Stanine", "Meaning", "", "Share of students", "", "Icon"], height=20)
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=4)
        ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=6)
        key_top = row + 1
        for k in range(9):
            r = key_top + k
            put(ws, f"B{r}", k + 1, f=font(10, True), al=CENTER, border=BOX)
            put(ws, f"C{r}", STANINE_WORDS[k], border=BOX)
            ws.merge_cells(f"C{r}:D{r}")
            put(ws, f"E{r}", f"about {EXPECTED_PCT[k]}%", al=CENTER, border=BOX)
            ws.merge_cells(f"E{r}:F{r}")
            put(ws, f"G{r}", k + 1, al=CENTER, border=BOX)
        ws.conditional_formatting.add(f"G{key_top}:G{key_top + 8}", traffic_lights(4, 7))
        row = key_top + 10
        notes = [
            "A stanine compares a student with their whole year group on that test: 5 is average, 1-3 below, "
            "7-9 above. Stanines from different tests, classes and year groups can be compared directly.",
            "Stanines settle once every class has entered its marks for a test; until then they are relative "
            "to the students entered so far.",
            "A whole year group always averages about stanine 5, so compare year groups using mean %. "
            "Compare classes using both.",
            "Avg KS2 is the mean of each student's KS2 maths and reading scaled scores (or the one that exists). "
            "Its band is a stanine within the year group, so 'vs KS2 band' shows progress from the starting point.",
            "This workbook holds personal data about children. Keep it in a staff-only location.",
        ]
        section(ws, f"B{row}", "Good to know")
        row += 1
        for text in notes:
            paragraph(ws, row, f"•   {text}")
            row += 1
        protect(ws)


# --------------------------------------------------------------- formatting

def traffic_lights(amber_from, green_from):
    return Rule(type="iconSet", iconSet=IconSet(
        iconSet="3TrafficLights1", showValue=None,
        cfvo=[FormatObject(type="percent", val=0), FormatObject(type="num", val=amber_from),
              FormatObject(type="num", val=green_from)]))


def arrows():
    return Rule(type="iconSet", iconSet=IconSet(
        iconSet="3Arrows", showValue=None,
        cfvo=[FormatObject(type="percent", val=0), FormatObject(type="num", val=-0.99),
              FormatObject(type="num", val=1)]))


def data_bar(color):
    return DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=color,
                       showValue=True)


def dxf_fill(color, bold=False):
    return DifferentialStyle(fill=PatternFill(bgColor=color, fill_type="solid"),
                             font=Font(bold=bold) if bold else None)


def add_diverging(ws, rng, first, mid, small, large):
    """Blue above / red below `mid` (a cell or number), in two steps each."""
    c = first
    steps = [(f"{c}-{mid}>={large}", BLUE_200), (f"{c}-{mid}>={small}", BLUE_100),
             (f"{c}-{mid}<=-{large}", RED_200), (f"{c}-{mid}<=-{small}", RED_100)]
    for cond, colour in steps:
        ws.conditional_formatting.add(rng, Rule(
            type="expression", dxf=dxf_fill(colour), stopIfTrue=True,
            formula=[f"AND(ISNUMBER({c}),ISNUMBER({mid}),{cond})"]))


def add_grid_rules(ws, rng, first, year_cell):
    c = first
    mean_pct = [(f"{c}-{year_cell}>=0.1", BLUE_200), (f"{c}-{year_cell}>=0.03", BLUE_100),
                (f"{c}-{year_cell}<=-0.1", RED_200), (f"{c}-{year_cell}<=-0.03", RED_100)]
    mean_stn = [(f"{c}-{year_cell}>=1", BLUE_200), (f"{c}-{year_cell}>=0.4", BLUE_100),
                (f"{c}-{year_cell}<=-1", RED_200), (f"{c}-{year_cell}<=-0.4", RED_100)]
    for measure, steps in (("Mean %", mean_pct), ("Mean stanine", mean_stn)):
        for cond, colour in steps:
            ws.conditional_formatting.add(rng, Rule(
                type="expression", dxf=dxf_fill(colour), stopIfTrue=True,
                formula=[f'AND(SelMeasure="{measure}",ISNUMBER({c}),ISNUMBER({year_cell}),{cond})']))
    for cond, colour in ((f"{c}=0", RED_200), (f"{c}<1", AMBER_100)):
        ws.conditional_formatting.add(rng, Rule(
            type="expression", dxf=dxf_fill(colour), stopIfTrue=True,
            formula=[f'AND(SelMeasure="Completion %",ISNUMBER({c}),{cond})']))


def add_number_format_rule(ws, rng, condition, fmt):
    ws.conditional_formatting.add(rng, Rule(
        type="expression", formula=[condition],
        dxf=DifferentialStyle(numFmt=NumberFormat(numFmtId=200, formatCode=fmt))))


def row_rule(key_cell):
    """A hairline under each row that has content, so empty slots stay clean."""
    return Rule(type="expression", formula=[f'{key_cell}<>""'], dxf=DifferentialStyle(
        border=Border(bottom=Side(style="thin", color=GRID))))


def completion_rules(first):
    return [Rule(type="expression", dxf=dxf_fill(RED_200), stopIfTrue=True,
                 formula=[f"AND(ISNUMBER({first}),{first}=0)"]),
            Rule(type="expression", dxf=dxf_fill(AMBER_100), stopIfTrue=True,
                 formula=[f"AND(ISNUMBER({first}),{first}<1)"])]


def completion_rules_counts(first):
    """'Not entered' counts: amber when some marks are missing for a class."""
    return [Rule(type="expression", dxf=dxf_fill(AMBER_100), stopIfTrue=True,
                 formula=[f"AND(ISNUMBER({first}),{first}>0)"])]


def paragraph(ws, row, text, first="B", last="H", chars_per_line=125):
    """Wrapped text across B:H, with the row tall enough for its lines."""
    put(ws, f"{first}{row}", text, al=WRAP)
    ws.merge_cells(f"{first}{row}:{last}{row}")
    ws.row_dimensions[row].height = 14.5 * max(1, math.ceil(len(text) / chars_per_line)) + 3


def dv_list(ws, ref, source):
    dv = DataValidation(type="list", formula1=source, allow_blank=False, showErrorMessage=True,
                        errorTitle="Pick from the list", error="Choose an option from the drop-down list.")
    ws.add_data_validation(dv)
    dv.add(ref)


def stat_tiles(ws, top, starts, ends, tiles, height_rows=4):
    for (label, formula, fmt), a, b in zip(tiles, starts, ends):
        put(ws, f"{a}{top}", label, f=font(9, True, INK2), fl=GREY_FILL, al=LEFT)
        ws.merge_cells(f"{a}{top}:{b}{top}")
        put(ws, f"{a}{top + 1}", formula, f=font(20, True, NAVY), fl=GREY_FILL, al=LEFT, fmt=fmt)
        ws.merge_cells(f"{a}{top + 1}:{b}{top + height_rows - 1}")
        for r in range(top, top + height_rows):
            for col in range(ws[f"{a}1"].column, ws[f"{b}1"].column + 1):
                ws.cell(r, col).fill = fill(GREY_FILL)


def print_setup(ws, fit_width):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    if fit_width:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0


def protect(ws):
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.selectLockedCells = False
    ws.protection.selectUnlockedCells = False


# ------------------------------------------------------------------- charts

def series(values_ref, cat=None, title_ref=None):
    s = Series(val=NumDataSource(numRef=NumRef(f=values_ref)), cat=cat)
    if title_ref:
        s.tx = SeriesLabel(strRef=StrRef(f=title_ref))
    return s


def _axis_text(size=900, color=INK2):
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(
        sz=size, solidFill=color, latin=DFont(typeface=FONT))), endParaRPr=CharacterProperties())])


def _style_chart(ch, title):
    ch.title = title
    ch.title.tx.rich.p[0].pPr = ParagraphProperties(defRPr=CharacterProperties(
        sz=1100, b=True, solidFill=INK, latin=DFont(typeface=FONT)))
    ch.style = 2
    ch.roundedCorners = False
    ch.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
    for ax in (ch.x_axis, ch.y_axis):
        ax.txPr = _axis_text()
        ax.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=BASE, w=9525))
        ax.delete = False
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=GRID, w=9525))


STATIC_CHARTS = False       # set by --static-charts, for LibreOffice renders only


def class_bar_chart():
    ch = BarChart()
    ch.type = "bar"
    ch.grouping = "clustered"
    _style_chart(ch, "Mean % by class, selected test")
    vals, cats = ("Calc!$AP$3:$AP$11", "Calc!$AO$3:$AO$11") if STATIC_CHARTS else ("Calc!chClassVals", "Calc!chClassCats")
    s = series(vals, cat=AxDataSource(strRef=StrRef(f=cats)))
    s.graphicalProperties = GraphicalProperties(solidFill=BLUE, ln=LineProperties(noFill=True))
    pt = DataPoint(idx=0)
    pt.graphicalProperties = GraphicalProperties(solidFill=MUTED, ln=LineProperties(noFill=True))
    s.dPt = [pt]
    s.dLbls = DataLabelList(showVal=True, showSerName=False, showCatName=False, showLegendKey=False,
                            numFmt="0%")
    s.dLbls.txPr = _axis_text(900, INK)
    ch.series.append(s)
    ch.gapWidth = 60
    ch.legend = None
    ch.x_axis.scaling.orientation = "maxMin"
    ch.y_axis.scaling.min = 0
    ch.y_axis.scaling.max = 1
    ch.y_axis.majorUnit = 0.25
    ch.y_axis.number_format = "0%"
    ch.y_axis.crosses = "max"
    ch.width, ch.height = 16, 7.2
    return ch


def stanine_chart():
    ch = BarChart()
    ch.type = "col"
    ch.grouping = "clustered"
    _style_chart(ch, "Stanine spread, selected test")
    for col, name, colour in (("AS", "This test", BLUE), ("AT", "Expected", "C3C2B7")):
        s = series(f"Calc!${col}$3:${col}$11", cat=AxDataSource(numRef=NumRef(f="Calc!$AR$3:$AR$11")),
                   title_ref=f"Calc!${col}$2")
        s.graphicalProperties = GraphicalProperties(solidFill=colour, ln=LineProperties(noFill=True))
        ch.series.append(s)
    ch.gapWidth = 50
    ch.overlap = -10
    ch.y_axis.number_format = "0%"
    ch.y_axis.scaling.min = 0
    ch.legend.position = "t"
    ch.legend.txPr = _axis_text()
    ch.width, ch.height = 16, 7.2
    return ch


def trend_chart():
    ch = LineChart()
    _style_chart(ch, "Focus class and whole year: mean stanine by test")
    refs = (("Calc!$AL$3:$AL$8", "Calc!$AM$3:$AM$8", "Calc!$AK$3:$AK$8") if STATIC_CHARTS else
            ("Calc!chTrendYear", "Calc!chTrendClass", "Calc!chTrendCats"))
    for ref, title_ref, colour, width in ((refs[0], "Calc!$AL$2", MUTED, 28575),
                                         (refs[1], "Dashboard!$C$7", BLUE, 25400)):
        s = series(ref, cat=AxDataSource(strRef=StrRef(f=refs[2])), title_ref=title_ref)
        s.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=colour, w=width))
        s.marker = Marker(symbol="circle", size=6)
        s.marker.graphicalProperties = GraphicalProperties(solidFill=colour, ln=LineProperties(solidFill=WHITE))
        s.smooth = False
        ch.series.append(s)
    ch.y_axis.scaling.min = 1
    ch.y_axis.scaling.max = 9
    ch.y_axis.majorUnit = 1
    ch.y_axis.number_format = "0"
    ch.legend.position = "t"
    ch.legend.txPr = _axis_text()
    ch.width, ch.height = 16, 9.5
    return ch


# ---------------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--students", help="the MIS 'All Students' export (.xlsx); omit for the fictional demo")
    ap.add_argument("--carry-over", action="append", default=[],
                    help="existing tracker whose tests and scores to bring across (repeatable)")
    ap.add_argument("--academic-year", default="2026-27")
    ap.add_argument("--school", default="Science")
    ap.add_argument("--out", default=str(HERE / "Mastery-Quiz-Tracker-demo.xlsm"))
    ap.add_argument("--no-cache", action="store_true",
                    help="skip filling in calculated values with LibreOffice (faster)")
    ap.add_argument("--static-charts", action="store_true", help=argparse.SUPPRESS)
    args = ap.parse_args()
    global STATIC_CHARTS
    STATIC_CHARTS = args.static_charts

    if args.students:
        students = load_student_export(args.students)
        tests = []
        for path in args.carry_over:
            tests += load_carry_over(path, students)
        demo = False
    else:
        students, tests = demo_data(args.academic_year)
        demo = True
    with_tests = [t.year for t in tests]
    busiest = max(YEARS, key=lambda y: (with_tests.count(y), y)) if tests else 7
    data = Dataset(students, tests, args.academic_year, args.school, demo,
                   dashboard_year=busiest, watch_year=busiest, teachers=demo_teachers() if demo else {})
    builder = Builder(data)
    wb = builder.build()
    out = Path(args.out)
    xlsm_package.save(wb, builder.extras, out, fill_cache=not args.no_cache)
    counts = {y: len(builder.students[y]) for y in YEARS}
    print(f"Wrote {out}  ({sum(counts.values())} students: "
          + ", ".join(f"Y{y} {n}" for y, n in counts.items()) + f"; {len(tests)} tests)")


if __name__ == "__main__":
    main()
