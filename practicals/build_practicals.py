#!/usr/bin/env python3
"""Build the Improved KS5 Practicals package, v1.1.

Inputs  : data/activities_v10.json      records extracted from the v1.0 documents
          data/source_register_v10.csv  the v1.0 file register
          data/implemented_changes_v10.csv
          v11_content.py                v1.1 fixes and newly written sheets
          v11_review.py                 red-team findings
Outputs : build/<PACKAGE>/...           the complete folder
          <PACKAGE>.zip                 next to this script

Run:  python3 build_practicals.py   then   python3 verify_practicals.py
"""
import csv, json, os, re, shutil, sys, zipfile, datetime, hashlib, html
from copy import deepcopy
from docx import Document
from docx.shared import Pt, Twips, RGBColor
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from v11_content import OVERRIDES, TECH_INTRO, ANSWERS_INTRO, SHARED
from v11_review import FINDINGS, DECISIONS

VERSION = "v1.1"
DATE_ISO = "2026-09-30"
DATE_TXT = "30 September 2026"
PACKAGE = f"Improved Practicals {DATE_ISO} {VERSION}"
STAMP = f"IMPROVED {VERSION} • {DATE_TXT}"
UNITS_AQA = True            # cm³ / mol dm⁻³ (False keeps the v1.0 mL / M)
SOURCE_ROOT_MARKER = "/06 - KS5 - Chemistry/"
BUILD = os.path.join(HERE, "build")
OUT = os.path.join(BUILD, PACKAGE)
ACT_DIR = "Activity resources"
SRC_DIR = "Source copies"

# ---------------------------------------------------------------- data ------
def units(s):
    if not UNITS_AQA or not isinstance(s, str):
        return s
    s = re.sub(r"\bmL\b", "cm³", s)
    s = re.sub(r"(\d)\s*M\b(?![\w(])", r"\1 mol dm⁻³", s)
    s = re.sub(r"\bg/L\b", "g dm⁻³", s)
    s = s.replace("per litre", "per dm³")
    s = s.replace("mol−1", "mol⁻¹")
    return s

def load_activities():
    raw = json.load(open(os.path.join(HERE, "data", "activities_v10.json"), encoding="utf-8"))
    acts = {int(k): v for k, v in raw.items()}
    for n, a in acts.items():
        a["answers"] = [tuple(x) for x in a["answers"]]
        a.setdefault("expected", ""); a.setdefault("hazards", ""); a.setdefault("v11", "")
        a["change_v10"] = a["change"]
    # merge overrides (two passes so COPY: references resolve)
    for n, ov in OVERRIDES.items():
        a = acts[n]
        for k, v in ov.items():
            if k == "prep_add":
                a["prep"] = a["prep"].rstrip() + " " + v
            else:
                a[k] = v
    for n, a in acts.items():
        for k, v in list(a.items()):
            if isinstance(v, str) and v.startswith("COPY:"):
                a[k] = deepcopy(acts[int(v[5:])][k])
    for n, a in acts.items():
        if not a["method"]:
            raise SystemExit(f"activity {n} has no method")
        if [q for q, _ in a["answers"]] != a["questions"]:
            raise SystemExit(f"activity {n}: questions and answers differ")
        for k in ("focus", "change", "skills", "equipment", "safety", "prep", "results", "expected", "time", "mode"):
            a[k] = units(a[k])
        a["method"] = [units(m) for m in a["method"]]
        a["questions"] = [units(q) for q in a["questions"]]
        a["answers"] = [(units(q), units(x)) for q, x in a["answers"]]
    return acts

def load_register():
    rows = list(csv.DictReader(open(os.path.join(HERE, "data", "source_register_v10.csv"), encoding="utf-8")))
    for r in rows:
        src = r["source"]
        i = src.find(SOURCE_ROOT_MARKER)
        r["source"] = src[i + len(SOURCE_ROOT_MARKER):] if i >= 0 else os.path.basename(src)
        r["activities"] = ";".join(x.strip() for x in re.split(r"[;,]", r["activities"]))
        r["replacement"] = r["replacement"].replace("IMPROVED v1.0", f"IMPROVED {VERSION}")
        r["canonical"] = ""      # filled after canonical names are known
    return rows

def load_changes():
    return {int(r["number"]): r for r in csv.DictReader(open(os.path.join(HERE, "data", "implemented_changes_v10.csv"), encoding="utf-8"))}

# ---------------------------------------------------------------- docx ------
def new_doc(title):
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Twips(11952), Twips(16848)
    s.top_margin = s.bottom_margin = Twips(936)
    s.left_margin = s.right_margin = Twips(1008)
    st = d.styles
    for name, size, bold, before, after in (("Normal", 10, None, None, 5), ("Heading 1", 13, True, 24, 0),
                                            ("Heading 2", 11, True, 12, 0), ("Title", 20, False, None, 15)):
        x = st[name]; x.font.name = "Arial"; x.font.size = Pt(size); x.font.color.rgb = RGBColor(0, 0, 0)
        if bold is not None: x.font.bold = bold
        rpr = x.element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
        if rf is None: rf = OxmlElement("w:rFonts"); rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), "Arial")
        if before is not None: x.paragraph_format.space_before = Pt(before)
        x.paragraph_format.space_after = Pt(after)
    # footer with page number
    fp = s.footer.paragraphs[0]
    r = fp.add_run(f"{STAMP}  |  "); r.font.size = Pt(8)
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
    rr = OxmlElement("w:r"); rpr = OxmlElement("w:rPr"); sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "16"); rpr.append(sz); rr.append(rpr)
    t = OxmlElement("w:t"); t.text = "1"; rr.append(t); fld.append(rr); fp._p.append(fld)
    z = d.settings.element.find(qn("w:zoom"))
    if z is not None and z.get(qn("w:percent")) is None: z.set(qn("w:percent"), "100")
    cp = d.core_properties
    cp.title = title; cp.author = "KS5 Chemistry practical review"; cp.comments = f"{STAMP}"
    now = datetime.datetime(2026, 9, 30, 12, 0, 0)
    cp.created = now; cp.modified = now; cp.last_modified_by = "build_practicals.py"; cp.revision = 2
    return d

def P(d, text, style=None):
    p = d.add_paragraph(text, style=style) if style else d.add_paragraph(text)
    return p
def H(d, text): return d.add_paragraph(text, style="Heading 1")
def H2(d, text): return d.add_paragraph(text, style="Heading 2")
def T(d, text): return d.add_paragraph(text, style="Title")
def page_break(d): d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
def lines(d, n=2):
    for _ in range(n): P(d, "_" * 68)

def table(d, header, rows, widths=None):
    t = d.add_table(rows=1, cols=len(header)); t.style = "Table Grid"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""; c.paragraphs[0].add_run(h).bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row): cells[i].text = str(v)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Twips(w)
    return t

def num(n): return f"{n:02d}"
def act_head(d, a, n, role):
    T(d, a["title"])
    P(d, f"Practical {num(n)} • {role} • {STAMP}")
    P(d, a["lesson"])
    P(d, f"{a['mode']} • {a['time']}")

def student_doc(a, n):
    d = new_doc(f"{num(n)} {a['title']} - Student")
    act_head(d, a, n, "Student")
    H(d, "Intended focus"); P(d, a["focus"])
    H(d, "Equipment and materials"); P(d, a["equipment"])
    H(d, "Working safely"); P(d, a["safety"])
    H(d, "Method")
    for i, m in enumerate(a["method"], 1): P(d, f"{i}. {m}")
    page_break(d)
    T(d, "Record and explain")
    P(d, f"Practical {num(n)} • Student • {STAMP}")
    P(d, "Name ____________________    Date ____________________")
    H(d, "Results")
    P(d, "Use a results table with these headings. Record observations before explanations.")
    P(d, a["results"])
    P(d, "Use your notebook or graph paper for results, calculations and graphs. Show units and working; keep reference or model data separate from measurements.")
    H(d, "Questions")
    for i, q in enumerate(a["questions"], 1):
        P(d, f"{i}. {q}"); lines(d)
    H(d, "What your conclusion should show")
    P(d, a["focus"])
    P(d, "State the evidence supporting your conclusion and one specific limitation. Where you collected data, explain the effect of that limitation rather than writing only “human error”.")
    lines(d, 3)
    return d

def teacher_doc(a, n):
    d = new_doc(f"{num(n)} {a['title']} - Teacher")
    act_head(d, a, n, "Teacher")
    H(d, "Intended focus"); P(d, a["focus"])
    H(d, "What changed"); P(d, a["change"])
    if a["v11"]: P(d, f"{VERSION}: {a['v11']}")
    H(d, "Skills"); P(d, a["skills"])
    H(d, "Equipment and quantities"); P(d, a["equipment"])
    H(d, "Controls and waste"); P(d, a["safety"])
    H(d, "Preparation and reliability"); P(d, a["prep"])
    if a["expected"]:
        H(d, "Expected results"); P(d, a["expected"])
    H(d, "Before first classroom use")
    P(d, "Document review completed. Local trial and current departmental risk assessment required for live procedures. Record actual timings, expected observations, reagent batches and any changes; update the paired student sheet when a method changes.")
    P(d, "Trial date __________  Technician __________  Teacher __________")
    P(d, "Outcome and amendments __________________________________________")
    page_break(d)
    T(d, "Teaching and answers")
    P(d, f"Practical {num(n)} • Teacher • {STAMP}")
    H(d, "Evidence of learning"); P(d, a["focus"])
    P(d, "Use the paired student sheet for the complete procedure. Observe the named skills directly; a demonstration or supplied dataset does not provide evidence that a student performed an experimental skill.")
    H(d, "Answers and expected reasoning")
    for q, x in a["answers"]:
        P(d, q); P(d, x)
    H(d, "Source and method status")
    P(d, "Rewritten from the Magnus Kerboodle resource set. The source register identifies every replaced file. Adaptations and replacement activities are stated on the first page. Expected results are predictions or reference values, not results from a school trial.")
    if a["sources"]:
        for s in a["sources"]: P(d, s)
    else:
        P(d, "Basis: original Kerboodle instructions and the chemical, stoichiometric and classroom review recorded in the change register. Use the departmental safety guidance for the actual stocks and apparatus.")
    H(d, "Lesson decision")
    P(d, f"{a['time']}. {a['mode']}. Retain time for the questions and conclusion; do not count handling equipment alone as meeting the learning focus.")
    return d

def guide_doc(acts):
    d = new_doc("Improved KS5 practical lesson guide")
    T(d, "Improved KS5 practical lesson guide"); P(d, STAMP)
    P(d, "Use this version with the improved student and teacher files. All 47 mapped entries have been revised and every entry now has a student sheet and a teacher sheet. This guide shows the intended learning, existing lesson match and realistic lesson commitment; the HTML index opens each paired resource.")
    P(d, "The method changes are adopted in this release. Original files remain untouched. Live procedures need a local technician trial and current departmental risk assessment before issue; no claim of physical testing is made. Sheets written in v1.1 for activities that had none in v1.0 are marked in the change log and need the same trial.")
    H(d, "Using the set")
    P(d, f"Open START HERE.html for the 47 pairs and the complete source register. Files marked Student are for learners; Teacher files include preparation, expected results, reliability checks and answers. The source copies folder preserves original names with an IMPROVED {VERSION} suffix. Use activity numbers to avoid mixing versions.")
    P(d, "Some formerly live activities now use models or reference data. They support conceptual learning, but cannot replace hands-on required practical evidence. " + SHARED)
    H(d, "What good practical teaching looks like")
    P(d, "Choose a specific learning purpose; explicitly teach the technique; control the variables that matter; give students a feasible measurement or observation; then require a conclusion supported by evidence. Build in controls, uncertainty and adequate time to interpret results. An entertaining demonstration alone is not evidence of individual practical competence.")
    P(d, "Reference: AQA A level Chemistry practical handbook, including practical competency and apparatus guidance. https://media.aqa.org.uk/resources/chemistry/AQA-7404-7405-PHBK.PDF")
    ns = sorted(acts)
    for i in range(0, len(ns), 3):
        grp = ns[i:i + 3]
        page_break(d)
        T(d, f"Practical lesson matches {grp[0]} to {grp[-1]}")
        for n in grp:
            a = acts[n]
            H(d, f"{num(n)} {a['title']}")
            rows = [("Topic", a["topic"]), ("Existing lesson", a["lesson"]), ("Intended focus", a["focus"]),
                    ("Skills", a["skills"]), ("Format and time", f"{a['mode']}; {a['time']}"), ("Implemented change", a["change"])]
            if a["v11"]: rows.append((f"Changed in {VERSION}", a["v11"]))
            table(d, ["Planning item", "Improved activity"], rows, widths=[2200, 7700])
    return d

def answers_doc(acts):
    d = new_doc("Improved practical answers")
    T(d, "Improved practical answers"); P(d, STAMP); P(d, ANSWERS_INTRO)
    for n in sorted(acts):
        a = acts[n]
        H(d, f"{num(n)} {a['title']}"); P(d, a["lesson"])
        for q, x in a["answers"]: P(d, q); P(d, x)
        if a["expected"]: P(d, f"Expected results: {a['expected']}")
    return d

def tech_doc(acts):
    d = new_doc("Improved practical technician preparation")
    T(d, "Improved practical technician preparation"); P(d, STAMP); P(d, TECH_INTRO)
    H(d, "Summary for booking and risk assessment")
    table(d, ["No", "Activity", "Mode", "Time", "Main hazards"],
          [(num(n), acts[n]["title"], acts[n]["mode"], acts[n]["time"], acts[n]["hazards"]) for n in sorted(acts)],
          widths=[500, 3000, 2000, 2000, 2400])
    for n in sorted(acts):
        a = acts[n]
        page_break(d) if n == 1 else None
        H(d, f"{num(n)} {a['title']}"); P(d, a["lesson"]); P(d, f"{a['mode']} • {a['time']}")
        H2(d, "Materials"); P(d, a["equipment"])
        H2(d, "Controls"); P(d, a["safety"])
        H2(d, "Preparation and trial"); P(d, a["prep"])
        if a["expected"]: H2(d, "Expected results"); P(d, a["expected"])
        P(d, "Trial date ________  Staff ________  Outcome ____________________")
    return d

def review_doc(inventory):
    d = new_doc(f"Red team review {VERSION}")
    T(d, f"Red team review of Improved Practicals v1.0 and changes made in {VERSION}"); P(d, STAMP)
    P(d, "Scope: document and chemistry review of the v1.0 package (all 69 shipped documents, both CSV registers and the HTML index). No procedure was carried out physically. External URLs could not be fetched from the review environment and remain unverified.")
    H(d, "Package inventory after rebuild")
    for k, v in inventory: P(d, f"{k}: {v}")
    H(d, "Findings and actions")
    table(d, ["ID", "Severity", "Area", "Activities", "Finding", f"Action in {VERSION}"],
          [f for f in FINDINGS], widths=[500, 900, 1100, 1100, 3200, 3100])
    H(d, "Decisions for the department")
    for i, x in enumerate(DECISIONS, 1): P(d, f"{i}. {x}")
    H(d, "Status")
    P(d, "All 47 activities have paired sheets, every index link resolves, and the verifier (verify_practicals.py) checks the inventory, links, units and a set of stoichiometric expectations on each build. Sheets written in v1.1 and the restored combustion practical have had document review only.")
    return d

# ---------------------------------------------------------------- build -----
def canon_name(n, title, role): return f"{num(n)} {title} - {role} - IMPROVED {VERSION}.docx"

def main():
    acts = load_activities(); reg = load_register(); chg = load_changes()
    if os.path.exists(OUT): shutil.rmtree(OUT)
    os.makedirs(os.path.join(OUT, ACT_DIR)); os.makedirs(os.path.join(OUT, SRC_DIR))
    docs = {}   # canonical relative path -> Document
    for n, a in acts.items():
        for role, fn in (("Student", student_doc), ("Teacher", teacher_doc)):
            rel = f"{ACT_DIR}/{canon_name(n, a['title'], role)}"
            docs[rel] = fn(a, n); docs[rel].save(os.path.join(OUT, rel))
    masters = {
        "guide": f"Improved KS5 practical lesson guide - IMPROVED {VERSION}.docx",
        "answers": f"Answers master - IMPROVED {VERSION}.docx",
        "tech": f"Technician preparation master - IMPROVED {VERSION}.docx",
        "review": f"Red team review - IMPROVED {VERSION}.docx",
    }
    guide_doc(acts).save(os.path.join(OUT, masters["guide"]))
    answers_doc(acts).save(os.path.join(OUT, masters["answers"]))
    tech_doc(acts).save(os.path.join(OUT, masters["tech"]))
    # source copies
    for r in reg:
        al = [int(x) for x in r["activities"].split(";")]
        dest = os.path.join(OUT, r["replacement"]); os.makedirs(os.path.dirname(dest), exist_ok=True)
        if r["role"] == "Technician":
            shutil.copy(os.path.join(OUT, masters["tech"]), dest); r["canonical"] = masters["tech"]
        elif r["role"] == "Answers":
            shutil.copy(os.path.join(OUT, masters["answers"]), dest); r["canonical"] = masters["answers"]
        elif len(al) == 1:
            n = al[0]; c = f"{ACT_DIR}/{canon_name(n, acts[n]['title'], r['role'])}"
            shutil.copy(os.path.join(OUT, c), dest); r["canonical"] = c
        else:   # one legacy file standing for several activities: all sheets in sequence
            fn = student_doc if r["role"] == "Student" else teacher_doc
            d = fn(acts[al[0]], al[0])
            for n in al[1:]:
                page_break(d); sub = fn(acts[n], n)
                for el in list(sub.element.body):
                    if el.tag != qn("w:sectPr"): d.element.body.insert(len(d.element.body) - 1, deepcopy(el))
            d.save(dest)
            r["canonical"] = "; ".join(f"{ACT_DIR}/{canon_name(n, acts[n]['title'], r['role'])}" for n in al)
    # registers
    fields = ["source", "source_sha256", "replacement", "activities", "role", "canonical"]
    with open(os.path.join(OUT, "Source register.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields); w.writeheader()
        for r in reg: w.writerow({k: r[k] for k in fields})
    json.dump({"version": VERSION, "date": DATE_ISO, "source_root": "…/02 - Resources/06 - KS5 - Chemistry/ (department OneDrive; see README)",
               "rows": [{k: r[k] for k in fields} for r in reg]},
              open(os.path.join(OUT, "Source register.json"), "w", encoding="utf-8"), indent=1, ensure_ascii=False)
    with open(os.path.join(OUT, "Implemented changes.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f); w.writerow(["number", "original_activity", "improved_activity", "implemented_change", f"change_in_{VERSION}", "intended_focus", "mode", "time", "verification"])
        for n in sorted(acts):
            a = acts[n]
            w.writerow([n, chg[n]["original_activity"], a["title"], a["change"], a["v11"] or "Units and metadata only", a["focus"], a["mode"], a["time"],
                        "Document review complete; live procedures require a local trial before first use."])
    with open(os.path.join(OUT, f"Changes v1.0 to {VERSION}.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f); w.writerow(["id", "severity", "area", "activities", "finding", "action"]); w.writerows(FINDINGS)
    # inventory, review, README, index
    n_act = len(os.listdir(os.path.join(OUT, ACT_DIR)))
    n_src = sum(len(fs) for _, _, fs in os.walk(os.path.join(OUT, SRC_DIR)))
    inventory = [("Activity resources", f"{n_act} files (47 student + 47 teacher)"), ("Source copies", f"{n_src} files for {len(reg)} register rows"),
                 ("Masters", "lesson guide, answers, technician preparation, red-team review"),
                 ("Sheets written in v1.1", ", ".join(num(n) for n in sorted(acts) if "written" in acts[n]["v11"]))]
    review_doc(inventory).save(os.path.join(OUT, masters["review"]))
    write_readme(acts, inventory, masters)
    write_index(acts, masters)
    # zip
    zpath = os.path.join(HERE, PACKAGE.replace(" ", "_") + ".zip")
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
        for dp, dn, fn in os.walk(OUT):
            for f in sorted(fn):
                full = os.path.join(dp, f); z.write(full, os.path.relpath(full, BUILD))
    print("built", OUT); print("zip", zpath, os.path.getsize(zpath), "bytes")

def write_readme(acts, inventory, masters):
    txt = f"""IMPROVED KS5 PRACTICALS {VERSION} — {DATE_TXT}

START HERE.html is the main index. It opens the lesson guide, the three masters, the red-team review and all 47 paired student and teacher resources.

Inventory (checked by verify_practicals.py on every build):
""" + "".join(f"  {k}: {v}\n" for k, v in inventory) + f"""
Activity resources: numbered sheets for teaching and preparation. Every activity has a Student and a Teacher sheet.
Source copies: separately named replacements for the 102 original files; the source register maps each one to its numbered activity. Where one original file stood for two activities (19 and 20) the copy contains both sheets. Legacy file names are kept on purpose so existing references still resolve; the content is the improved activity named inside.
Source register: paths are relative to the department folder '…/02 - Resources/06 - KS5 - Chemistry/'. The absolute OneDrive prefix present in v1.0 was removed because it identified the author.
Answers master: answers and expected results for the revised questions in this set.
Technician preparation master: summary table, then quantities, controls, trial notes and expected results for every activity.
Red team review: findings on v1.0 and the action taken in {VERSION}; 'Changes v1.0 to {VERSION}.csv' is the same list as data.
Implemented changes.csv: v1.0 change, {VERSION} change and intended focus for every activity.

Units: volumes and concentrations use cm³ and mol dm⁻³ to match AQA examinations.

Shared methods: {SHARED}

Sheets written in {VERSION} for activities that had none in v1.0 (12, 19, 20, 25, 26, 27, 28, 36, 38, 40, 41, 42, 43, 44, 45, 46) and the restored combustion practical (17) have had document review only.

All originals remain unchanged. These files have undergone document and chemistry review, not physical trials. Teachers and technicians should trial live procedures with the actual equipment and reagents and apply current departmental risk assessments before student use. Model and reference-data activities are labelled and do not substitute for hands-on required-practical competence evidence.
"""
    open(os.path.join(OUT, "README.txt"), "w", encoding="utf-8").write(txt)

def write_index(acts, masters):
    e = html.escape
    rows = []
    for n in sorted(acts):
        a = acts[n]
        s = f"{ACT_DIR}/{canon_name(n, a['title'], 'Student')}"; t = f"{ACT_DIR}/{canon_name(n, a['title'], 'Teacher')}"
        flag = f" <span class=new>{e(VERSION)}</span>" if "written" in a["v11"] or n == 17 else ""
        rows.append(f"<tr><td>{num(n)}</td><td>{e(a['title'])}{flag}</td><td><a href=\"{e(s)}\">Student</a> <a href=\"{e(t)}\">Teacher</a></td><td>{e(a['lesson'])}</td><td>{e(a['time'])}</td></tr>")
    body = f"""<!DOCTYPE html><html lang="en-GB"><head><meta charset="utf-8"><title>Improved KS5 practicals {e(VERSION)}</title>
<style>body{{font:16px Arial;max-width:1200px;margin:40px auto;color:#18364c}}table{{border-collapse:collapse;width:100%}}td,th{{padding:10px;border:1px solid #ccc;text-align:left;vertical-align:top}}tr:nth-child(even){{background:#f0f3f5}}a{{margin-right:12px}}h1,h2{{color:#111}}.new{{font-size:12px;background:#ffe9a8;padding:2px 5px;border-radius:3px}}</style></head><body>
<h1>Improved KS5 practicals</h1><p>{e(STAMP)}</p>
<p>47 activities, each with a Student and a Teacher sheet. Corrected methods, intended focus, teacher preparation, expected results and matching answers. Originals preserved. Live procedures require local trial and departmental risk assessment before use. Entries marked <span class=new>{e(VERSION)}</span> were written or rewritten in this version and have had document review only.</p>
<p><a href="{e(masters['guide'])}">Lesson planning guide</a> <a href="{e(masters['tech'])}">Technician master</a> <a href="{e(masters['answers'])}">Answers master</a> <a href="{e(masters['review'])}">Red team review</a> <a href="Changes v1.0 to {e(VERSION)}.csv">Change list (CSV)</a> <a href="Source register.csv">Source register</a> <a href="README.txt">README</a></p>
<p>Shared methods: {e(SHARED)}</p>
<table><tr><th>No</th><th>Activity</th><th>Resources</th><th>Lesson</th><th>Time</th></tr>
{''.join(rows)}
</table></body></html>"""
    open(os.path.join(OUT, "START HERE.html"), "w", encoding="utf-8").write(body)

if __name__ == "__main__":
    main()
