#!/usr/bin/env python3
"""Verify the built package: inventory, links, register, units, duplicates and
a set of stoichiometric expectations.  Exit code 1 on any failure."""
import csv, json, os, re, sys, html
import docx
from build_practicals import OUT, ACT_DIR, SRC_DIR, VERSION, UNITS_AQA, load_activities, canon_name, num, PACKAGE, HERE

fails = []
def check(cond, msg):
    if not cond: fails.append(msg)

def text_of(path):
    d = docx.Document(path); out = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows: out.extend(c.text for c in r.cells)
    for s in d.sections: out.extend(p.text for p in s.footer.paragraphs)
    return "\n".join(out)

acts = load_activities()
check(os.path.isdir(OUT), f"missing build dir {OUT}")

# 1. inventory
act_files = sorted(os.listdir(os.path.join(OUT, ACT_DIR)))
check(len(act_files) == 94, f"expected 94 activity files, found {len(act_files)}")
for n, a in acts.items():
    for role in ("Student", "Teacher"):
        check(canon_name(n, a["title"], role) in act_files, f"missing canonical {n} {role}")

# 2. index links
idx = open(os.path.join(OUT, "START HERE.html"), encoding="utf-8").read()
links = [html.unescape(l) for l in re.findall(r'href="([^"]+)"', idx)]
check(len(links) >= 94 + 7, f"index has only {len(links)} links")
for l in links:
    if not l.startswith("http"): check(os.path.exists(os.path.join(OUT, l)), f"broken index link: {l}")
check(idx.count("<tr>") == 48, "index should have 47 activity rows")

# 3. register
rows = list(csv.DictReader(open(os.path.join(OUT, "Source register.csv"), encoding="utf-8")))
check(len(rows) == 102, f"register rows {len(rows)} != 102")
for r in rows:
    check(os.path.exists(os.path.join(OUT, r["replacement"])), f"register replacement missing: {r['replacement']}")
    for c in r["canonical"].split("; "): check(os.path.exists(os.path.join(OUT, c)), f"register canonical missing: {c}")
    check("," not in r["activities"], f"comma separator in activities: {r['activities']}")
    check(not r["source"].startswith("/"), f"absolute source path leaked: {r['source']}")
    check(r["replacement"].endswith(f"IMPROVED {VERSION}.docx"), f"bad replacement name {r['replacement']}")
js = json.load(open(os.path.join(OUT, "Source register.json"), encoding="utf-8"))
check(len(js["rows"]) == 102 and js["version"] == VERSION, "register json mismatch")

# 4. every docx: opens, no v1.0 stamp, no personal path, units
all_docx = []
for dp, dn, fn in os.walk(OUT):
    for f in fn:
        if f.endswith(".docx"): all_docx.append(os.path.join(dp, f))
check(len(all_docx) == 94 + 102 + 4, f"docx count {len(all_docx)} != 200")
texts = {}
for p in all_docx:
    try:
        t = text_of(p)
    except Exception as e:
        fails.append(f"cannot open {p}: {e}"); continue
    texts[p] = t
    rel = os.path.relpath(p, OUT)
    check("IMPROVED v1.0" not in t, f"old stamp in {rel}")
    check(f"IMPROVED {VERSION}" in t, f"no {VERSION} stamp in {rel}")
    check("jamesdavies" not in t and ("OneDrive" not in t or "Red team review" in rel), f"personal path in {rel}")
    if UNITS_AQA and "Red team review" not in rel:
        check(not re.search(r"\bmL\b", t), f"mL remains in {rel}")
        m = re.search(r"\d\s*M\b(?![\w(])", t)
        check(not m, f"M remains in {rel}: {t[max(0,m.start()-30):m.end()+20] if m else ''}")
        check("mol−1" not in t, f"mol−1 remains in {rel}")
for f in ("README.txt", "Source register.csv", "Source register.json", "Implemented changes.csv"):
    t = open(os.path.join(OUT, f), encoding="utf-8").read()
    check("jamesdavies" not in t and "OneDrive-" not in t, f"personal path in {f}")

# 5. sheet structure
def sheet(n, role): return texts[os.path.join(OUT, ACT_DIR, canon_name(n, acts[n]["title"], role))]
for n, a in acts.items():
    s, t = sheet(n, "Student"), sheet(n, "Teacher")
    check(s.startswith(a["title"]), f"{n} student title")
    check(f"Practical {num(n)} • Student • IMPROVED {VERSION}" in s, f"{n} student meta")
    check(f"Practical {num(n)} • Teacher • IMPROVED {VERSION}" in t, f"{n} teacher meta")
    for i in range(1, len(a["method"]) + 1): check(f"\n{i}. " in s, f"{n} method step {i} missing")
    check(len(a["method"]) >= 5, f"{n} method has only {len(a['method'])} steps")
    for q, x in a["answers"]:
        check(q in s and q in t and x in t, f"{n} question/answer not present in both sheets: {q[:40]}")
    check(a["expected"] == "" or a["expected"] in t, f"{n} expected results missing from teacher sheet")
    check(bool(a["hazards"]), f"{n} has no hazards entry for the technician table")

# 6. deliberate duplicates and the 17/18 split
same = lambda x, y: acts[x]["method"] == acts[y]["method"]
for x, y in ((22, 23), (27, 28), (31, 35), (26, 47)): check(same(x, y), f"{x}/{y} should share a method")
check(not same(17, 18), "17 and 18 must differ")

# 7. stoichiometric expectations (value must appear in the teacher sheet text)
Mr = dict(succinic=118.09, KHP=204.22, CaCO3=100.09, NaCl=58.44, NaHCO3=84.01, Na2CO3=105.99, PbI2=461.0,
          Mohr=392.14, glucose=180.16, EtOH=46.07, thio=248.18, salicylic=138.12, aspirin=180.16, MgSO4=120.37, hept=246.47, cyclohexanol=100.16, cyclohexene=82.14)
def expect(n, value, fmt, label):
    s = fmt.format(value); check(s in sheet(n, "Teacher"), f"{n} {label}: expected '{s}' in teacher sheet")
expect(4, 0.00125 / (2.000 / Mr["succinic"] / 0.2500) * 1000, "{:.2f}", "titre")             # 18.45
expect(5, 0.0025 / 0.07350 * 1000, "{:.2f}", "titre")                                            # 34.01
expect(7, (0.02500 - 2 * 0.800 / Mr["CaCO3"]) / 10 / 0.1000 * 1000, "{:.2f}", "titre")          # 9.01
expect(8, 0.0100 * Mr["NaCl"], "{:.3f}", "NaCl mass")                                            # 0.584
expect(9, 2.000 / Mr["NaHCO3"] / 2 * Mr["Na2CO3"], "{:.3f}", "residue")                          # 1.262
expect(9, 2.000 - 2.000 / Mr["NaHCO3"] / 2 * Mr["Na2CO3"], "{:.3f}", "loss")                     # 0.738
expect(10, 0.00250 * Mr["PbI2"], "{:.3f}", "plateau")                                            # 1.153
expect(16, 0.1000 * Mr["Mohr"], "{:.2f}", "Mohr salt g per dm3")                                 # 39.21
expect(19, 3.50 / Mr["NaHCO3"], "{:.4f}", "NaHCO3 mol")                                          # 0.0417
expect(19, 2 * 2.50 / Mr["Na2CO3"], "{:.4f}", "HCl mol for Na2CO3")                              # 0.0472
expect(20, 0.0250 * Mr["MgSO4"], "{:.3f}", "anhydrous mass")                                     # 3.009
expect(20, 0.0250 * Mr["hept"], "{:.3f}", "hydrate mass")                                        # 6.162
expect(20, 50.00 - 0.0250 * 7 * 18.015, "{:.2f}", "added water")                                 # 46.85
expect(32, 2 * 5.00 / Mr["glucose"] * Mr["EtOH"], "{:.2f}", "ethanol mass")                      # 2.56
expect(36, 0.0100 * Mr["thio"], "{:.3f}", "thiosulfate g per dm3")                               # 2.482
expect(37, 0.050 * 120, "{:.1f}", "charge")                                                      # 6.0
expect(40, 0.0500 * Mr["Na2CO3"], "{:.4f}", "carbonate g per dm3")                               # 5.2995
expect(45, 6.00 * Mr["aspirin"] / Mr["salicylic"], "{:.2f}", "aspirin theoretical")             # 7.83
expect(18, 0.0250 * 57000 / (50 * 4.18), "{:.1f}", "ΔT")                                        # 6.8
check("16.5" in sheet(4, "Teacher") and "20.5" in sheet(4, "Teacher"), "4 corrected titre range")
check("hexane" in sheet(12, "Student"), "12 caveat about non-polar liquids")
check("1.0 mol dm⁻³ Na2CO3" in sheet(46, "Student"), "46 carbonate concentration")
check("nominally" in sheet(16, "Student"), "16 concentration wording")

# 8. zip present and complete
z = os.path.join(HERE, PACKAGE.replace(" ", "_") + ".zip")
check(os.path.exists(z), "zip missing")
if os.path.exists(z):
    import zipfile
    names = set(zipfile.ZipFile(z).namelist())
    for dp, dn, fn in os.walk(OUT):
        for f in fn:
            rel = os.path.relpath(os.path.join(dp, f), os.path.dirname(OUT))
            check(rel in names, f"not in zip: {rel}")

if fails:
    print(f"FAIL ({len(fails)})"); [print(" -", f) for f in fails[:60]]; sys.exit(1)
print(f"OK: 94 activity files, {len(rows)} source copies, {len(all_docx)} docx, {len(links)} index links, {len(acts)} activities checked")
