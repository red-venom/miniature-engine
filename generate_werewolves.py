#!/usr/bin/env python3
"""Generate Werewolves.xlsx — a playable, formula-driven Werewolves (Mafia) game.

The workbook is the game engine: it deals secret roles from a seed, resolves
night attacks against doctor saves, answers the Seer, tallies day votes and
declares the winner. The moderator only ever types into yellow cells.

Usage:  python3 generate_werewolves.py [--out Werewolves.xlsx]

Requires: openpyxl (pip install openpyxl). No macros — pure formulas, so the
file works in Excel, Google Sheets and LibreOffice.
"""

import argparse

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

# ---------------------------------------------------------------- layout maps
N_PLAYERS = 20          # player slots
N_ROUNDS = 10           # supported nights/days

SETUP_TOP = 14          # Setup: first player row (rows 14..33)
SETUP_BOT = SETUP_TOP + N_PLAYERS - 1

NIGHT_TOP = 4           # Night: rows 4..13 hold nights 1..10
NIGHT_BOT = NIGHT_TOP + N_ROUNDS - 1

DAY_VOTE_TOP = 4        # Day: vote matrix rows 4..23
DAY_VOTE_BOT = DAY_VOTE_TOP + N_PLAYERS - 1
DAY_TALLY_TOP = DAY_VOTE_BOT + 3            # tally rows 26..45
DAY_TALLY_BOT = DAY_TALLY_TOP + N_PLAYERS - 1
DAY_LEADER = DAY_TALLY_BOT + 2              # 47: most votes
DAY_TIE = DAY_LEADER + 1                    # 48: tie flag
DAY_AUTO = DAY_LEADER + 2                   # 49: voted out (auto)
DAY_OVERRIDE = DAY_LEADER + 3               # 50: moderator override
DAY_FINAL = DAY_LEADER + 4                  # 51: eliminated today

BOARD_TOP = 11          # Game Board: player rows 11..30
BOARD_BOT = BOARD_TOP + N_PLAYERS - 1

NAMES = f"Setup!$C${SETUP_TOP}:$C${SETUP_BOT}"          # player names
ROLETAB = f"Setup!$C${SETUP_TOP}:$F${SETUP_BOT}"        # name -> role (col 4)
PACK = f"Setup!$H${SETUP_TOP}:$H${SETUP_BOT}"           # werewolf names
VICTIMS = f"Night!$G${NIGHT_TOP}:$G${NIGHT_BOT}"        # night kills
LYNCHED = f"Day!$C${DAY_FINAL}:$L${DAY_FINAL}"          # day eliminations
ORDTAB = f"'Game Board'!$B${BOARD_TOP}:$F${BOARD_BOT}"  # name -> death phase (col 5)
BOARD_NAMES = f"'Game Board'!$B${BOARD_TOP}:$B${BOARD_BOT}"
BOARD_STATUS = f"'Game Board'!$D${BOARD_TOP}:$D${BOARD_BOT}"

SAMPLE_PLAYERS = ["Aria", "Ben", "Chloe", "Dev", "Elle", "Finn", "Grace", "Hugo"]

ROLES = [
    ("Werewolf", "🐺", "Each night you and your pack silently agree on one villager to eat. "
                       "By day: act shocked, point fingers, survive the vote. You win when the "
                       "wolves are as many as the rest of the town."),
    ("Seer", "🔮", "Each night, choose one player — the moderator tells you whether they are a "
                   "werewolf. Deadly knowledge: reveal too much and the wolves will come for you."),
    ("Doctor", "💉", "Each night, choose one player to protect (yourself included). If the wolves "
                     "attack that player, they survive the night."),
    ("Hunter", "🏹", "You die loudly. Whether eaten or voted out, you immediately fire one last "
                     "shot and take a player of your choice down with you."),
    ("Villager", "🌾", "No powers — just eyes, ears and a vote. Find the liars, rally the town and "
                       "vote the wolves out before you're dinner."),
]

# ------------------------------------------------------------------- styling
INK = "1F2937"          # near-black text
MUTE = "9CA3AF"         # grey helper text
HDR_BG = "2F3A4F"       # slate header
INPUT_BG = "FFF3BF"     # yellow: "type here"
CALC_BG = "F3F4F6"      # grey: automatic
WOLF_RED = "B3261E"
GOOD_GREEN = "1E7D32"
WARN_AMBER = "B45309"
NIGHT_BG = "22304A"
BANNER_BG = "E8EBF0"

F_TITLE = Font(size=16, bold=True, color=INK)
F_SUB = Font(size=10, italic=True, color="6B7280")
F_HDR = Font(size=10, bold=True, color="FFFFFF")
F_BODY = Font(size=10, color=INK)
F_BOLD = Font(size=10, bold=True, color=INK)
F_MUTE = Font(size=8, color=MUTE)
F_BIG = Font(size=20, bold=True, color=INK)

FILL_HDR = PatternFill("solid", fgColor=HDR_BG)
FILL_IN = PatternFill("solid", fgColor=INPUT_BG)
FILL_CALC = PatternFill("solid", fgColor=CALC_BG)
FILL_BANNER = PatternFill("solid", fgColor=BANNER_BG)

THIN = Side(style="thin", color="C3C8D0")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)


def put(ws, ref, value=None, font=F_BODY, fill=None, align=LEFT, border=None):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = font
    if fill is not None:
        c.fill = fill
    c.alignment = align
    if border is not None:
        c.border = border
    return c


def header_row(ws, row, cols_titles, height=22):
    ws.row_dimensions[row].height = height
    for col, title in cols_titles:
        put(ws, f"{col}{row}", title, font=F_HDR, fill=FILL_HDR, align=CENTER, border=BOX)


def input_cell(ws, ref, value=None):
    return put(ws, ref, value, font=F_BOLD, fill=FILL_IN, align=CENTER, border=BOX)


def calc_cell(ws, ref, formula, font=F_BODY, align=LEFT):
    return put(ws, ref, formula, font=font, fill=FILL_CALC, align=align, border=BOX)


def name_dropdown(ws, ranges):
    dv = DataValidation(type="list", formula1="=PlayerNames", allow_blank=True)
    dv.error = "Pick a player from the list (names come from the Setup sheet)."
    dv.errorTitle = "Unknown player"
    ws.add_data_validation(dv)
    for r in ranges:
        dv.add(r)


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


# ------------------------------------------------------------------- sheets
def build_howto(ws):
    ws.sheet_properties.tabColor = "7F8C8D"
    widths(ws, {"A": 2, "B": 110})
    lines = [
        ("🐺 WEREWOLVES — the spreadsheet edition", F_TITLE),
        ("A parlour game of lies, fangs and formulas. 5–20 players plus one moderator, who drives this workbook.", F_SUB),
        ("", None),
        ("THE IDEA", F_HDR),
        ("A few of you are secretly werewolves; everyone else is a villager (some with special powers).", None),
        ("Each NIGHT the wolves secretly eat someone. Each DAY the whole village votes someone out.", None),
        ("The village wins when every wolf is gone. The wolves win when they are as many as the rest of the town.", None),
        ("", None),
        ("COLOUR CODE", F_HDR),
        ("🟡 Yellow cells are the only cells you ever type in. Grey cells compute themselves — leave them be.", None),
        ("", None),
        ("HOW TO RUN A GAME — the moderator's checklist", F_HDR),
        ("1.  SETUP sheet: type the player names, set the role counts (it suggests a wolf count) and any seed number.", None),
        ("2.  ROLE CARDS sheet: pass the device around — each player picks their name, memorises their secret role,", None),
        ("     clears the yellow cell, and passes it on. (Suspicious table? The moderator whispers roles instead.)", None),
        ("3.  NIGHT sheet: “Everybody close your eyes.” Wake in order — Werewolves choose a victim, the Doctor", None),
        ("     protects someone, the Seer inspects someone. Type the three names into that night's row.", None),
        ("4.  MORNING: read the 🌅 announcement aloud. The Game Board updates the body count automatically.", None),
        ("5.  DAY sheet: debate! Then record each living player's vote in today's column — the tally below names", None),
        ("     the eliminated player. On a tie nobody dies (or use the moderator override row after a revote).", None),
        ("6.  Repeat night and day. The GAME BOARD banner declares the winner the instant a side has won.", None),
        ("7.  Hunter in play? The moment the Hunter dies they shoot one player — record it on the Game Board.", None),
        ("", None),
        ("THE ROLES", F_HDR),
    ]
    lines += [(f"{emoji}  {name} — {desc}", None) for name, emoji, desc in ROLES]
    lines += [
        ("", None),
        ("THE MODERATOR'S NIGHT SCRIPT (steal these lines)", F_HDR),
        ("“Village, close your eyes.”  ·  “Werewolves — open your eyes. Choose your prey.”", None),
        ("“Werewolves, sleep. Doctor — who do you save tonight?”  ·  “Doctor, sleep. Seer — point at a soul,", None),
        ("and I shall tell you what walks inside it.”  ·  “Everyone, wake up. I'm afraid there's been… an incident.”", None),
        ("", None),
        ("FINE PRINT", F_HDR),
        ("•  Player names must be unique (the formulas look players up by name).", None),
        ("•  Works in Excel 2019+, Google Sheets and LibreOffice. In older Excel only the wolf-pack line on Role Cards may show #NAME?.", None),
        ("•  Up to 20 players and 10 rounds. Change the seed on Setup to re-deal roles — never mid-game!", None),
        ("•  The sheet trusts the moderator: it flags mistakes (targeting the dead, ties) but never blocks a dramatic twist.", None),
        ("•  Honour system: players seeing Role Cards must resist peeking at Setup or the Game Board. Wolves have no honour — everyone else does.", None),
    ]
    for i, (text, font) in enumerate(lines, start=2):
        c = ws.cell(row=i, column=2, value=text or None)
        c.font = font or F_BODY
        if font is F_HDR:
            c.fill = FILL_HDR
    ws.sheet_view.showGridLines = False


def build_setup(ws):
    ws.sheet_properties.tabColor = "F1C40F"
    widths(ws, {"A": 2, "B": 16, "C": 18, "D": 13, "E": 10, "F": 22, "G": 2, "H": 18})
    put(ws, "B1", "🐺 Werewolves — Setup", F_TITLE)
    put(ws, "B2", "🟡 yellow = you type · grey = automatic. Fill names, tweak roles, then deal with the Role Cards sheet.", F_SUB)

    put(ws, "B4", "Game seed", F_BOLD)
    input_cell(ws, "C4", 42)
    put(ws, "D4", "any whole number — change it to re-deal every role", F_MUTE)

    n_players = f"COUNTA($C${SETUP_TOP}:$C${SETUP_BOT})"
    config = [
        ("B6", "🐺 Werewolves", "C6", 2, f'="suggested: "&MAX(1,ROUNDDOWN({n_players}/4,0))'),
        ("B7", "🔮 Seer", "C7", 1, "0 or 1"),
        ("B8", "💉 Doctor", "C8", 1, "0 or 1"),
        ("B9", "🏹 Hunter", "C9", 0, "0 or 1"),
    ]
    for lref, label, cref, default, note in config:
        put(ws, lref, label, F_BOLD)
        input_cell(ws, cref, default)
        put(ws, f"D{lref[1:]}", note, F_MUTE)
    put(ws, "B10", "🌾 Villagers", F_BOLD)
    calc_cell(ws, "C10", f"=IF({n_players}=0,0,{n_players}-SUM($C$6:$C$9))", align=CENTER)
    put(ws, "D10", "automatic: everyone else", F_MUTE)

    ws.merge_cells("B12:F12")
    put(ws, "B12",
        f'=IF({n_players}=0,"",IF($C$10<0,"⚠️ More roles than players — reduce the special roles!",'
        f'IF($C$6<1,"⚠️ You need at least one Werewolf.",'
        f'IF($C$6>={n_players}-$C$6,"⚠️ Too many wolves — they would win instantly!",'
        f'"✅ Setup looks good — deal the roles!"))))', F_BOLD)

    header_row(ws, SETUP_TOP - 1, [("B", "#"), ("C", "Player name"), ("D", "🎲 shuffle"),
                                   ("E", "deal order"), ("F", "🤫 SECRET ROLE"), ("H", "🐺 pack (mod only)")])
    for i in range(N_PLAYERS):
        r = SETUP_TOP + i
        put(ws, f"B{r}", i + 1, F_MUTE, align=CENTER, border=BOX)
        input_cell(ws, f"C{r}", SAMPLE_PLAYERS[i] if i < len(SAMPLE_PLAYERS) else None)
        calc_cell(ws, f"D{r}",
                  f'=IF($C{r}="","",ROUND(ABS(SIN($C$4*1.98987+ROW()*7.23319))*100000,4)+ROW()/100000)',
                  font=F_MUTE, align=CENTER)
        calc_cell(ws, f"E{r}", f'=IF($C{r}="","",RANK(D{r},$D${SETUP_TOP}:$D${SETUP_BOT},1))', align=CENTER)
        calc_cell(ws, f"F{r}",
                  f'=IF($C{r}="","",IF($C$10<0,"⚠️",'
                  f'IF(E{r}<=$C$6,"Werewolf",IF(E{r}<=$C$6+$C$7,"Seer",'
                  f'IF(E{r}<=$C$6+$C$7+$C$8,"Doctor",IF(E{r}<=SUM($C$6:$C$9),"Hunter","Villager"))))))',
                  font=F_BOLD, align=CENTER)
        calc_cell(ws, f"H{r}", f'=IF($F{r}="Werewolf",$C{r},"")', font=F_MUTE, align=CENTER)

    role_rng = f"F{SETUP_TOP}:F{SETUP_BOT}"
    ws.conditional_formatting.add(role_rng, FormulaRule(
        formula=[f'$F{SETUP_TOP}="Werewolf"'], font=Font(bold=True, color=WOLF_RED)))
    ws.conditional_formatting.add("B12", FormulaRule(
        formula=['LEFT($B$12,2)="⚠️"'], font=Font(bold=True, color=WOLF_RED)))
    ws.freeze_panes = f"A{SETUP_TOP}"
    ws.sheet_view.showGridLines = False


def build_role_cards(ws):
    ws.sheet_properties.tabColor = "8E44AD"
    widths(ws, {"A": 2, "B": 14, "C": 18, "D": 14, "E": 14, "F": 14, "G": 14, "H": 14,
                "I": 2, "J": 12, "K": 5, "L": 46})
    put(ws, "B1", "🎭 Role Cards — for players' eyes only", F_TITLE)

    put(ws, "B3", "1) Pick your name:", F_BOLD)
    input_cell(ws, "C3")
    name_dropdown(ws, ["C3"])

    ws.merge_cells("B5:H6")
    put(ws, "B5",
        '=IF($C$3="","👆 pick your name above…",'
        'IFERROR(VLOOKUP(VLOOKUP($C$3,' + ROLETAB + ',4,FALSE),$J$3:$L$7,2,FALSE)&"  You are: "&'
        'UPPER(VLOOKUP($C$3,' + ROLETAB + ',4,FALSE)),"(roles not dealt yet)"))',
        font=F_BIG, fill=FILL_CALC, align=CENTER, border=BOX)

    ws.merge_cells("B8:H11")
    put(ws, "B8",
        '=IF($C$3="","",IFERROR(VLOOKUP(VLOOKUP($C$3,' + ROLETAB + ',4,FALSE),$J$3:$L$7,3,FALSE),""))',
        font=Font(size=12, color=INK), fill=FILL_CALC, align=WRAP, border=BOX)

    ws.merge_cells("B13:H13")
    put(ws, "B13",
        '=IF($C$3="","",IF(IFERROR(VLOOKUP($C$3,' + ROLETAB + ',4,FALSE),"")="Werewolf",'
        '"🐺 Your pack: "&TEXTJOIN(", ",TRUE,' + PACK + '),""))',
        font=Font(size=12, bold=True, color=WOLF_RED), fill=FILL_CALC, align=CENTER, border=BOX)

    ws.merge_cells("B15:H16")
    put(ws, "B15", "2) Memorise it, CLEAR your name (select the yellow cell, press Delete), "
                   "then pass the device on. No scrolling to other sheets — you're on your honour!",
        font=Font(size=11, bold=True, color=WARN_AMBER), align=WRAP)

    put(ws, "J2", "cheat sheet (public info)", F_MUTE)
    for i, (name, emoji, desc) in enumerate(ROLES):
        r = 3 + i
        put(ws, f"J{r}", name, F_MUTE)
        put(ws, f"K{r}", emoji, F_MUTE, align=CENTER)
        put(ws, f"L{r}", desc, F_MUTE)
    ws.sheet_view.showGridLines = False


def build_night(ws):
    ws.sheet_properties.tabColor = NIGHT_BG
    widths(ws, {"A": 7, "B": 18, "C": 18, "D": 18, "E": 30, "F": 48, "G": 14, "H": 36})
    put(ws, "A1", "🌙 Night — moderator only", F_TITLE)
    put(ws, "A2", "Wake the roles in order: Werewolves → Doctor → Seer. Fill in one row per night; the sheet resolves it.", F_SUB)
    header_row(ws, NIGHT_TOP - 1, [("A", "Night"), ("B", "🐺 wolves attack"), ("C", "💉 doctor protects"),
                                   ("D", "🔮 seer inspects"), ("E", "🔮 the seer learns…"),
                                   ("F", "🌅 morning announcement"), ("G", "victim (auto)"), ("H", "⚠️ checks")])

    def dead_before(cell, night_cell):
        return (f'IF(AND({cell}<>"",IFERROR(VLOOKUP({cell},{ORDTAB},5,FALSE),999)<2*{night_cell}-1,'
                f'IFERROR(VLOOKUP({cell},{ORDTAB},5,FALSE),999)<>""),"⚠️ "&{cell}&" is already dead. ","")')

    for i in range(N_ROUNDS):
        r = NIGHT_TOP + i
        put(ws, f"A{r}", i + 1, F_BOLD, align=CENTER, border=BOX)
        for col in "BCD":
            input_cell(ws, f"{col}{r}")
        calc_cell(ws, f"E{r}",
                  f'=IF($D{r}="","",$D{r}&" → "&IFERROR(IF(VLOOKUP($D{r},{ROLETAB},4,FALSE)="Werewolf",'
                  f'"🐺 IS A WEREWOLF","✅ not a werewolf"),"?"))')
        calc_cell(ws, f"F{r}",
                  f'=IF($B{r}="","",IF($B{r}=$C{r},'
                  f'"🛡️ The wolves attacked "&$B{r}&" — but the Doctor was there! Nobody died.",'
                  f'"☠️ "&$B{r}&" was killed in the night."))')
        calc_cell(ws, f"G{r}", f'=IF(OR($B{r}="",$B{r}=$C{r}),"",$B{r})', font=F_MUTE, align=CENTER)
        calc_cell(ws, f"H{r}",
                  '=TRIM(' + "&".join(dead_before(f"${c}{r}", f"$A{r}") for c in "BCD") + ')',
                  font=Font(size=9, color=WARN_AMBER))
    name_dropdown(ws, [f"B{NIGHT_TOP}:D{NIGHT_BOT}"])

    rng = f"F{NIGHT_TOP}:F{NIGHT_BOT}"
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'ISNUMBER(SEARCH("killed",$F{NIGHT_TOP}))'], font=Font(color=WOLF_RED, bold=True)))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'ISNUMBER(SEARCH("Doctor",$F{NIGHT_TOP}))'], font=Font(color=GOOD_GREEN, bold=True)))
    ws.freeze_panes = f"A{NIGHT_TOP}"
    ws.sheet_view.showGridLines = False


def build_day(ws):
    ws.sheet_properties.tabColor = "E67E22"
    widths(ws, {"A": 4, "B": 20} | {get_column_letter(3 + d): 13 for d in range(N_ROUNDS)})
    put(ws, "B1", "☀️ Day — the village votes", F_TITLE)
    put(ws, "B2", "In each living player's row, enter who they vote to eliminate today. 💀 rows are dead — they don't vote.", F_SUB)

    header_row(ws, 3, [("A", ""), ("B", "Player")] +
               [(get_column_letter(3 + d), f"Day {d + 1}") for d in range(N_ROUNDS)])
    for i in range(N_PLAYERS):
        r = DAY_VOTE_TOP + i
        s = SETUP_TOP + i
        calc_cell(ws, f"A{r}",
                  f'=IF($B{r}="","",IF(COUNTIFS({BOARD_NAMES},$B{r},{BOARD_STATUS},"DEAD")>0,"💀",""))',
                  align=CENTER)
        calc_cell(ws, f"B{r}", f'=IF(Setup!$C{s}="","",Setup!$C{s})', font=F_BOLD)
        for d in range(N_ROUNDS):
            input_cell(ws, f"{get_column_letter(3 + d)}{r}")

    put(ws, "B{}".format(DAY_TALLY_TOP - 1), "🗳 votes received", F_HDR, fill=FILL_HDR, border=BOX)
    for i in range(N_PLAYERS):
        r = DAY_TALLY_TOP + i
        v = DAY_VOTE_TOP + i
        calc_cell(ws, f"B{r}", f'=IF($B{v}="","",$B{v})', font=F_MUTE)
        for d in range(N_ROUNDS):
            col = get_column_letter(3 + d)
            calc_cell(ws, f"{col}{r}",
                      f'=IF($B{r}="","",COUNTIF({col}${DAY_VOTE_TOP}:{col}${DAY_VOTE_BOT},$B{r}))',
                      font=F_MUTE, align=CENTER)

    tally = f"${DAY_TALLY_TOP}:{{c}}${DAY_TALLY_BOT}"
    rows = [
        (DAY_LEADER, "Most votes",
         '=IF(COUNT({c}' + tally + ')=0,"",IF(MAX({c}' + tally + ')=0,"",'
         f'INDEX($B${DAY_TALLY_TOP}:$B${DAY_TALLY_BOT},MATCH(MAX({{c}}{tally}),{{c}}{tally},0))))'),
        (DAY_TIE, "Tie?",
         '=IF({c}$' + str(DAY_LEADER) + '="","",IF(COUNTIF({c}' + tally + ',MAX({c}' + tally + '))>1,"⚠️ TIE — revote or nobody",""))'),
        (DAY_AUTO, "Voted out (auto)",
         '=IF({c}$' + str(DAY_LEADER) + '="","",IF({c}$' + str(DAY_TIE) + '<>"","",{c}$' + str(DAY_LEADER) + '))'),
        (DAY_OVERRIDE, "Moderator override", None),
        (DAY_FINAL, "☠️ ELIMINATED TODAY",
         '=IF({c}$' + str(DAY_OVERRIDE) + '<>"",{c}$' + str(DAY_OVERRIDE) + ',{c}$' + str(DAY_AUTO) + ')'),
    ]
    for row, label, template in rows:
        put(ws, f"B{row}", label, F_BOLD, border=BOX)
        for d in range(N_ROUNDS):
            col = get_column_letter(3 + d)
            if template is None:
                input_cell(ws, f"{col}{row}")
            else:
                calc_cell(ws, f"{col}{row}", template.format(c=col), align=CENTER,
                          font=F_BOLD if row == DAY_FINAL else F_BODY)

    name_dropdown(ws, [f"C{DAY_VOTE_TOP}:{get_column_letter(2 + N_ROUNDS)}{DAY_VOTE_BOT}",
                       f"C{DAY_OVERRIDE}:{get_column_letter(2 + N_ROUNDS)}{DAY_OVERRIDE}"])

    last = get_column_letter(2 + N_ROUNDS)
    ws.conditional_formatting.add(f"A{DAY_VOTE_TOP}:{last}{DAY_VOTE_BOT}", FormulaRule(
        formula=[f'$A{DAY_VOTE_TOP}="💀"'], fill=PatternFill("solid", fgColor="E5E7EB"),
        font=Font(color=MUTE, strike=True)))
    ws.conditional_formatting.add(f"C{DAY_TIE}:{last}{DAY_TIE}", FormulaRule(
        formula=[f'C{DAY_TIE}<>""'], font=Font(bold=True, color=WARN_AMBER)))
    ws.conditional_formatting.add(f"C{DAY_FINAL}:{last}{DAY_FINAL}", FormulaRule(
        formula=[f'C{DAY_FINAL}<>""'], font=Font(bold=True, color=WOLF_RED)))
    ws.conditional_formatting.add(f"C{DAY_TALLY_TOP}:{last}{DAY_TALLY_BOT}", FormulaRule(
        formula=[f'AND($B{DAY_TALLY_TOP}<>"",C{DAY_TALLY_TOP}<>"",C{DAY_TALLY_TOP}>0,'
                 f'C{DAY_TALLY_TOP}=MAX(C${DAY_TALLY_TOP}:C${DAY_TALLY_BOT}))'],
        fill=PatternFill("solid", fgColor="FADBD8")))
    ws.freeze_panes = f"C{DAY_VOTE_TOP}"
    ws.sheet_view.showGridLines = False


def build_board(ws):
    ws.sheet_properties.tabColor = "27AE60"
    widths(ws, {"A": 4, "B": 20, "C": 12, "D": 9, "E": 34, "F": 13, "G": 18, "H": 2, "I": 16, "J": 10})
    put(ws, "B1", "🐺 Werewolves — Moderator's Board", F_TITLE)

    ws.merge_cells("B3:G3")
    ws.row_dimensions[3].height = 30
    put(ws, "B3",
        f'=IF(COUNTA({NAMES})=0,"👋 Add players on the Setup sheet to begin",'
        f'IF($E$5=0,"🎉 THE VILLAGE WINS — every last werewolf is gone!",'
        f'IF($E$5>=$G$5,"🐺 THE WEREWOLVES WIN — the pack overruns the village!",'
        f'"⚔️ Game on — "&$E$5&" werewolf(s) hiding among "&$G$5&" townsfolk")))',
        font=Font(size=14, bold=True, color=INK), fill=FILL_BANNER, align=CENTER, border=BOX)

    counts = [("B5", "🙂 Alive", "C5", f'=COUNTIF({BOARD_STATUS},"ALIVE")'),
              ("D5", "🐺 Wolves", "E5",
               f'=COUNTIFS($C${BOARD_TOP}:$C${BOARD_BOT},"Werewolf",{BOARD_STATUS},"ALIVE")'),
              ("F5", "🏘 Town", "G5", "=$C$5-$E$5")]
    for lref, label, cref, formula in counts:
        put(ws, lref, label, F_BOLD)
        calc_cell(ws, cref, formula, font=F_BOLD, align=CENTER)

    ws.merge_cells("B7:F7")
    put(ws, "B7", "🏹 Hunter's revenge — the moment the Hunter dies, they shoot one player. Record it here:", F_BOLD)
    input_cell(ws, "G7")
    name_dropdown(ws, ["G7"])

    put(ws, "I4", "Hunter (auto)", F_MUTE)
    calc_cell(ws, "J4",
              f'=IFERROR(INDEX({NAMES},MATCH("Hunter",Setup!$F${SETUP_TOP}:$F${SETUP_BOT},0)),"")',
              font=F_MUTE, align=CENTER)
    put(ws, "I5", "fell in phase #", F_MUTE)
    calc_cell(ws, "J5",
              f'=IF($J$4="",999,MIN(IFERROR(MATCH($J$4,{VICTIMS},0)*2-1,999),'
              f'IFERROR(MATCH($J$4,{LYNCHED},0)*2,999)))',
              font=F_MUTE, align=CENTER)

    header_row(ws, BOARD_TOP - 1, [("A", "#"), ("B", "Player"), ("C", "Role"),
                                   ("D", "Status"), ("E", "Fate"), ("F", "phase # (auto)")])
    for i in range(N_PLAYERS):
        r = BOARD_TOP + i
        s = SETUP_TOP + i
        ordinal = (f'MIN(IFERROR(MATCH($B{r},{VICTIMS},0)*2-1,999),'
                   f'IFERROR(MATCH($B{r},{LYNCHED},0)*2,999),'
                   f'IF(AND($G$7<>"",$B{r}=$G$7),$J$5,999))')
        put(ws, f"A{r}", i + 1, F_MUTE, align=CENTER, border=BOX)
        calc_cell(ws, f"B{r}", f'=IF(Setup!$C{s}="","",Setup!$C{s})', font=F_BOLD)
        calc_cell(ws, f"C{r}", f'=IF($B{r}="","",Setup!$F{s})', align=CENTER)
        calc_cell(ws, f"D{r}", f'=IF($B{r}="","",IF($F{r}="","ALIVE","DEAD"))', font=F_BOLD, align=CENTER)
        calc_cell(ws, f"E{r}",
                  f'=IF($F{r}="","",IF(AND($G$7<>"",$B{r}=$G$7,$F{r}=$J$5),"🏹 Shot by the Hunter\'s last arrow",'
                  f'IF(MOD($F{r},2)=1,"☠️ Night "&($F{r}+1)/2&" — eaten by wolves",'
                  f'"🗳 Day "&$F{r}/2&" — voted out")))')
        calc_cell(ws, f"F{r}", f'=IF($B{r}="","",IF({ordinal}=999,"",{ordinal}))', font=F_MUTE, align=CENTER)

    ws.conditional_formatting.add("B3", FormulaRule(
        formula=['ISNUMBER(SEARCH("VILLAGE WINS",$B$3))'],
        fill=PatternFill("solid", fgColor="D5F5DC"), font=Font(size=14, bold=True, color=GOOD_GREEN)))
    ws.conditional_formatting.add("B3", FormulaRule(
        formula=['ISNUMBER(SEARCH("WEREWOLVES WIN",$B$3))'],
        fill=PatternFill("solid", fgColor="FADBD8"), font=Font(size=14, bold=True, color=WOLF_RED)))
    body = f"A{BOARD_TOP}:F{BOARD_BOT}"
    ws.conditional_formatting.add(body, FormulaRule(
        formula=[f'$D{BOARD_TOP}="DEAD"'], fill=PatternFill("solid", fgColor="E5E7EB"),
        font=Font(color=MUTE, strike=True)))
    ws.conditional_formatting.add(f"C{BOARD_TOP}:C{BOARD_BOT}", FormulaRule(
        formula=[f'$C{BOARD_TOP}="Werewolf"'], font=Font(bold=True, color=WOLF_RED)))
    ws.conditional_formatting.add(f"D{BOARD_TOP}:D{BOARD_BOT}", FormulaRule(
        formula=[f'$D{BOARD_TOP}="ALIVE"'], font=Font(bold=True, color=GOOD_GREEN)))
    ws.freeze_panes = f"A{BOARD_TOP}"
    ws.sheet_view.showGridLines = False


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--out", default="Werewolves.xlsx")
    args = ap.parse_args()

    wb = Workbook()
    wb.properties.title = "Werewolves — the spreadsheet edition"
    wb.properties.creator = "generate_werewolves.py"

    howto = wb.active
    howto.title = "How to Play"
    setup = wb.create_sheet("Setup")
    cards = wb.create_sheet("Role Cards")
    night = wb.create_sheet("Night")
    day = wb.create_sheet("Day")
    board = wb.create_sheet("Game Board")

    wb.defined_names["PlayerNames"] = DefinedName("PlayerNames", attr_text=NAMES)

    build_howto(howto)
    build_setup(setup)
    build_role_cards(cards)
    build_night(night)
    build_day(day)
    build_board(board)

    wb.save(args.out)
    print(f"wrote {args.out}")


if __name__ == "__main__":
    main()
